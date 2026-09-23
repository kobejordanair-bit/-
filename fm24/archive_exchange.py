"""Lossless workbook / HTML exchange. Shared by CLI and the browser worker.

The workbook is preserved byte for byte. It is read, never rewritten. Imported
HTML is parsed as data; none of its scripts are executed or copied as code.
"""
from __future__ import annotations

import argparse
import base64
from contextlib import closing
import datetime as dt
import hashlib
import io
import json
from pathlib import Path
import re
import sqlite3
import tempfile
import zipfile
import xml.etree.ElementTree as ET

from etl import load, normalise, slug_table

FORMAT = 'FM24_EXCHANGE_V1'
ROOT = Path(__file__).parent
FILES = ['etl.py', 'build_site.py', 'resolver.py', 'data_rules.py', 'evidence.py',
         'history.py', 'experience.py', 'player_awards.py', 'honour_review.py',
         'comparison.py', 'integrity.py', 'table_roles.py', 'archive_exchange.py',
         'template.html', 'experience.js', 'experience.css', 'honour_features.js',
         'player_compare.js', 'archive_workflow.js', 'archive_worker.js']
MAX_INPUT = 80 * 1024 * 1024
MAX_EXPANDED = 300 * 1024 * 1024


def dumps(value):
    return json.dumps(value, ensure_ascii=False, separators=(',', ':')).replace('<', '\\u003c')


def digest(value):
    return hashlib.sha256(value if isinstance(value, bytes) else value.encode('utf-8')).hexdigest()


def runtime():
    files = {name: (ROOT / name).read_text(encoding='utf-8') for name in FILES}
    return {'engine': digest(dumps(files)), 'files': files}


def validate_workbook(raw: bytes):
    if len(raw) > MAX_INPUT:
        raise ValueError('檔案超過 80 MB，請使用本機工具處理。')
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if len(z.infolist()) > 10000 or sum(i.file_size for i in z.infolist()) > MAX_EXPANDED:
            raise ValueError('工作簿解壓縮規模超過上限。')
        if 'xl/workbook.xml' not in z.namelist():
            raise ValueError('不是有效的 .xlsx 工作簿。')
        if any('vbaProject' in n for n in z.namelist()):
            raise ValueError('請將含巨集的工作簿另存為 .xlsx 後匯入。')
        missing_cache = []
        errors = []
        ns = {'s': 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'}
        for name in z.namelist():
            if not re.fullmatch(r'xl/worksheets/sheet\d+\.xml', name):
                continue
            for _, cell in ET.iterparse(io.BytesIO(z.read(name)), events=['end']):
                if cell.tag != '{%s}c' % ns['s']:
                    continue
                if cell.find('s:f', ns) is not None and cell.find('s:v', ns) is None:
                    missing_cache.append(f'{name}:{cell.get("r")}')
                elif cell.find('s:f', ns) is not None:
                    v = cell.find('s:v', ns)
                    if not v.text and cell.get('t') != 'str':
                        missing_cache.append(f'{name}:{cell.get("r")}')
                if cell.get('t') == 'e':
                    errors.append(f'{name}:{cell.get("r")}')
                cell.clear()
        if missing_cache or errors:
            raise ValueError('Excel 有未快取的公式或錯誤儲存格，請先在 Excel 重新計算並儲存：' +
                             ', '.join((missing_cache + errors)[:8]))


def manifest(con):
    sheets = []
    for name, table, count, headers in con.execute('SELECT * FROM _sheets ORDER BY sheet_name'):
        rows = con.execute(f'SELECT * FROM "{table}" ORDER BY _row').fetchall()
        hashes = [digest(dumps(list(r[:-1]))) for r in rows]
        sheets.append({'name': name, 'rows': count, 'columns': headers.split('\t'),
                       'sha256': digest(dumps([headers, rows])), 'rowHashes': hashes})
    return {'sheets': sheets, 'rows': sum(s['rows'] for s in sheets)}


def verify_mirror(workbook, con):
    import openpyxl
    wb = openpyxl.load_workbook(workbook, read_only=True, data_only=True)
    try:
        if set(wb.sheetnames) != {r[0] for r in con.execute('SELECT sheet_name FROM _sheets')}:
            raise ValueError('Excel 與資料庫工作表不同，請使用完整更新流程重新建置。')
        for ws in wb:
            it = ws.iter_rows(values_only=True)
            header = next(it, ())
            table = slug_table(ws.title)
            actual = con.execute(f'SELECT * FROM "{table}" ORDER BY _row')
            if '\t'.join('' if h is None else str(h) for h in header) != con.execute('SELECT columns FROM _sheets WHERE sheet_name=?', (ws.title,)).fetchone()[0]:
                raise ValueError('Excel 與資料庫欄位不符：' + ws.title)
            for i, row in enumerate(it, 2):
                values = [normalise(v) for v in row[:len(header)]]
                values += [None] * (len(header)-len(values))
                if all(v is None for v in values): continue
                expected = tuple(None if v is None else str(v) for v in values) + (i,)
                if actual.fetchone() != expected:
                    raise ValueError(f'Excel 與資料庫不是同一版：{ws.title} 第 {i} 列')
            if actual.fetchone() is not None:
                raise ValueError('資料庫有 Excel 沒有的資料列：' + ws.title)
    finally:
        wb.close()


def make_exchange(workbook: Path, db: Path, payload: dict, engine: str, *, verify=True):
    raw = workbook.read_bytes()
    with closing(sqlite3.connect(db)) as con:
        if verify:
            validate_workbook(raw)
            verify_mirror(workbook, con)
        inventory = manifest(con)
    payload_json = dumps(payload)
    return {'format': FORMAT, 'engine': engine, 'created': dt.datetime.now(dt.timezone.utc).isoformat(),
            'filename': workbook.name, 'workbook': base64.b64encode(raw).decode('ascii'),
            'workbookSha256': digest(raw), 'manifest': inventory,
            'payloadJSON': payload_json, 'payloadSha256': digest(payload_json)}


def convert(workbook: Path, engine: str):
    """One isolated build; its failure cannot change any current archive."""
    from build_site import Archive, collect
    validate_workbook(workbook.read_bytes())
    with tempfile.TemporaryDirectory(prefix='fm24-update-') as directory:
        db = Path(directory) / 'archive.sqlite'
        load(workbook, db)
        a = Archive(db)
        try:
            names = {r['table_name'] for r in a.q('SELECT table_name FROM _sheets')}
            required = {'Player_Dim', 'Club_Dim', 'Barcelona_Season_Master', 'Workbook_Schema_Metadata'}
            if required - names:
                raise ValueError('不是完整 World Master；缺少：' + ', '.join(sorted(required - names)))
            payload = collect(a)
            if not payload['seasons'] or not payload['world']['seasons']:
                raise ValueError('缺少可展示的賽季，已保留目前版本。')
            return make_exchange(workbook, db, payload, engine, verify=False)
        finally:
            a.con.close()


def render_html(packet, code, *, portable=True, standalone=True):
    from build_site import STANDALONE_SKELETON
    html = code['files']['template.html']
    for marker, filename in [('/*__EXPERIENCE_JS__*/', 'experience.js'),
                             ('/*__EXPERIENCE_CSS__*/', 'experience.css'),
                             ('/*__HONOUR_JS__*/', 'honour_features.js'),
                             ('/*__COMPARISON_JS__*/', 'player_compare.js'),
                             ('/*__WORKFLOW_JS__*/', 'archive_workflow.js')]:
        html = html.replace(marker, code['files'][filename], 1)
    metadata = {k: v for k, v in packet.items() if k != 'payloadJSON'}
    metadata['portable'] = portable
    html = html.replace('"__ARCHIVE_DATA__"', packet['payloadJSON'], 1)
    html = html.replace('"__ARCHIVE_EXCHANGE__"', dumps(metadata), 1)
    html = html.replace('"__ARCHIVE_RUNTIME__"', dumps(code), 1)
    return STANDALONE_SKELETON.format(body=html) if standalone else html


def extract(text):
    """JSON decoder only; HTML scripts never run."""
    if len(text.encode('utf-8')) > MAX_INPUT:
        raise ValueError('輸入超過 80 MB。')
    if text.lstrip().startswith('{'):
        packet = json.loads(text)
    else:
        m = re.search(r'<script\b[^>]*\bid=["\']fm24-exchange["\'][^>]*>(.*?)</script\s*>', text, re.S | re.I)
        if not m:
            raise ValueError('這份舊 HTML 未保存原始 Excel；請另選原始 .xlsx。')
        packet = json.loads(m.group(1))
        app = re.search(r'<script\b[^>]*\bid=["\']fm24-app["\'][^>]*>(.*?)</script\s*>', text, re.S | re.I)
        source = app.group(1) if app else text
        start = re.search(r'\b(?:let|const) DATA = ', source)
        if not start:
            raise ValueError('找不到網站資料。')
        value, end = json.JSONDecoder().raw_decode(source[start.end():])
        packet['payloadJSON'] = source[start.end():start.end() + end]
    if packet.get('format') != FORMAT:
        raise ValueError('不是 FM24 完整交換包。')
    raw = base64.b64decode(packet['workbook'], validate=True)
    if digest(raw) != packet['workbookSha256'] or digest(packet['payloadJSON']) != packet['payloadSha256']:
        raise ValueError('檔案雜湊核對失敗，資料可能已損壞。')
    validate_workbook(raw)
    return packet, raw


def main():
    parser = argparse.ArgumentParser(description='Excel / FM24 HTML / 交換 JSON → 完整新版資料夾（不覆寫來源）')
    parser.add_argument('input', type=Path)
    parser.add_argument('-o', '--output', type=Path, required=True, help='新的輸出資料夾；存在時拒絕覆寫')
    parser.add_argument('--site', action='store_true', help='產生網站基準版；預設為可攜版')
    args = parser.parse_args()
    if args.output.exists():
        parser.error('輸出資料夾已存在，請指定新版本名稱。')
    code = runtime()
    with tempfile.TemporaryDirectory() as d:
        workbook = Path(d) / 'World_Master.xlsx'
        if args.input.suffix.lower() == '.xlsx':
            raw = args.input.read_bytes()
        else:
            _, raw = extract(args.input.read_text(encoding='utf-8-sig'))
        workbook.write_bytes(raw)
        packet = convert(workbook, code['engine'])
        packet['filename'] = args.input.name if args.input.suffix.lower() == '.xlsx' else 'World_Master.xlsx'
        args.output.parent.mkdir(parents=True, exist_ok=True)
        with tempfile.TemporaryDirectory(prefix='.fm24-ready-', dir=args.output.parent) as stage:
            dest = Path(stage) / 'version'
            dest.mkdir()
            (dest / 'archive.html').write_text(render_html(packet, code, portable=not args.site), encoding='utf-8')
            (dest / 'World_Master.xlsx').write_bytes(raw)
            (dest / 'fm24-exchange.json').write_text(dumps(packet), encoding='utf-8')
            (dest / 'fm24-data.json').write_text(packet['payloadJSON'], encoding='utf-8')
            checksums = {p.name: digest(p.read_bytes()) for p in dest.iterdir()}
            (dest / 'SHA256SUMS.json').write_text(dumps(checksums), encoding='utf-8')
            dest.rename(args.output)
    print(f'完成：{args.output}；{len(packet["manifest"]["sheets"])} 表 / {packet["manifest"]["rows"]:,} 列')


if __name__ == '__main__':
    main()
