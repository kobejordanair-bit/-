"""Real importer/site acceptance on an explicit isolated copy of the current Master.

Creates one clearly synthetic test player. Never publish these outputs as save data.
Does not claim to test Hermes image recognition or a remote installation.
"""
import argparse
import hashlib
import json
from pathlib import Path
import sys
import traceback

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pipeline as p


def accept(workbook, output):
    from openpyxl import load_workbook
    before = p.sha(workbook)
    p.configure(output, workbook=workbook)
    cfg = p.config_at(output / 'config.json')
    # Recreate a known existing source profile from current Master only.
    wb = load_workbook(workbook, read_only=True)
    try:
        it = wb['Player_Profile_Snapshots'].iter_rows(values_only=True)
        columns = next(it); r = dict(zip(columns, next(it)))
        known_dobs = {dict(zip(columns, row)).get('DOB') for row in it}
    finally:
        wb.close()
    mapping = {'Name': 'Player_Name_Raw', 'Raw_Name': 'Player_Name_Raw', 'Snapshot': 'Snapshot_Date',
               'Nationality': 'Nationality_Raw', 'Date_of_Birth': 'DOB', 'Current_Club': 'Current_Club_Raw',
               'Position': 'Position_Raw', 'International_Caps': 'Senior_NT_Caps', 'International_Goals': 'Senior_NT_Goals',
               'Shirt_Number': 'Shirt_Number', 'U21_Caps': 'U21_Caps', 'U21_Goals': 'U21_Goals'}
    known = 'FORMAT=FM_WORLD_IMPORT_V1_1\n\nPLAYER\n' + '\n'.join(k + '=' + str(r[v]) for k, v in mapping.items() if r.get(v) is not None) + '\n'

    def batch(text, name):
        source = output / name; source.write_text(text, encoding='utf-8')
        b = p.prepare(cfg, [source], kind='transcript')
        m = p.read_json(b['batch']); m['documents'][0]['review']['checked'] = True
        p.write_json(b['batch'], m)
        return b['batch']

    results = {'mode': 'ISOLATED_SYNTHETIC_ACCEPTANCE_NOT_GAME_FACTS', 'original_sha256': before}
    print('1/4 Existing real profile replay; expect no changes', flush=True)
    replay = p.run_batch(cfg, batch(known, 'known-profile.txt'))
    assert replay['status'] == 'NO_OP', replay
    results['known_profile_replay'] = 'NO_OP'
    dob = next(f'1990-01-{d:02d}' for d in range(1, 29) if f'1990-01-{d:02d}' not in known_dobs)
    synthetic = ('FORMAT=FM_WORLD_IMPORT_V1_1\n\nPLAYER\nName=FM24_PIPELINE_TEST_ONLY\n'
                 f'Date_of_Birth={dob}\nSnapshot=2035-06-02\nCurrent_Club=巴塞隆納\n'
                 '\nCLUB_TOTALS\nSeason|Club|Apps_Raw|Goals|Assists|Rating\n'
                 '2034/35|巴塞隆納|3(2)|3|NULL|7.25\n')
    b = batch(synthetic, 'SYNTHETIC-TEST-ONLY.txt')
    print('2/4 Real importer + website + activation, synthetic fixture only', flush=True)
    run = p.run_batch(cfg, b)
    assert run['status'] == 'ACTIVE', run
    results['activation'] = run
    pid = run['steps'][0]['player_id']
    release = Path(run['release'])
    payload = p.read_json(release / 'fm24-data.json')
    row = payload['comparison']['players'][pid]['club'][0]
    assert row['apps'] == 5 and row['goals'] == 3 and row['assists'] is None, row
    assert row['adoption'] == 'ADOPTED' and row['fact'].startswith('FACT-'), row
    from archive_exchange import extract
    packet, raw = extract((release / 'archive.html').read_text(encoding='utf-8'))
    assert hashlib.sha256(raw).hexdigest() == p.sha(release / 'World_Master.xlsx')
    results['website_visible_values'] = {k: row[k] for k in ('apps', 'goals', 'assists', 'adoption', 'fact')}
    # The current curated personal award fields must remain unchanged.
    old = load_workbook(workbook, read_only=True); new = load_workbook(release / 'World_Master.xlsx', read_only=True)
    try:
        for sheet in ('Canonical_Award_Facts', 'Award_Fact_Source_Links'):
            assert list(old[sheet].values) == list(new[sheet].values), sheet
        def keyed(w):
            rows = w['Barcelona_Player_Career'].iter_rows(values_only=True); cols = next(rows)
            return {r[0]: dict(zip(cols, r)) for r in rows}
        a, z = keyed(old), keyed(new)
        assert z[pid]['Assists'] is None, 'Unknown assists must not become zero in Excel career totals'
        for ident in a:
            for field in ('Major_Awards', 'FIFPro_Count', 'Ballon_dOr_Top3_Count', 'Derivation_Sources'):
                assert a[ident][field] == z[ident][field], (ident, field)
    finally:
        old.close(); new.close()
    print('3/4 Replay twice: same receipt, then new transcription whitespace', flush=True)
    active_sha = p.active_registry(cfg['registry'])['sha256']
    assert p.run_batch(cfg, b)['status'] == 'NO_OP'
    assert p.run_batch(cfg, batch(synthetic + '\n', 'SEMANTIC-REPLAY.txt'))['status'] == 'NO_OP'
    assert p.active_registry(cfg['registry'])['sha256'] == active_sha
    print('4/4 Rollback and original-file preservation', flush=True)
    p.rollback(cfg, active_sha)
    assert p.active_registry(cfg['registry'])['sha256'] == before
    assert p.sha(workbook) == before
    results.update(pass_=True, **{'pass': True}, original_unchanged=True, rollback=True,
                   same_batch_idempotent=True, semantic_replay_idempotent=True,
                   original_awards_preserved=True, html_excel_byte_identical=True,
                   live_hermes_vision_tested=False)
    p.write_json(output / 'acceptance.json', results)
    print(json.dumps({k: v for k, v in results.items() if k != 'activation'}, ensure_ascii=False, indent=2))


if __name__ == '__main__':
    ap = argparse.ArgumentParser(); ap.add_argument('--workbook', type=Path, required=True); ap.add_argument('--output', type=Path, required=True)
    a = ap.parse_args(); accept(a.workbook.resolve(), a.output.resolve())
