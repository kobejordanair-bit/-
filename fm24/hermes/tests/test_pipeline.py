from contextlib import ExitStack
import json
from pathlib import Path
import sys
import tempfile
import unittest
from unittest.mock import patch

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
import pipeline as p
from import_policy import select_row, evidence_same
from screenshot_provenance import stamp, FIELDS, adopted
from generic_importer import parse_apps

TEXT = 'FORMAT=FM_WORLD_IMPORT_V1_1\n\nPLAYER\nName=Fixture Player\nSnapshot=2035-03-01\n'
STATS = '\nCLUB_TOTALS\nSeason|Club|Apps_Raw|Goals|Assists|Rating\n2034/35|Fixture Club|10(3)|8|0|7.25\n'


class ReceiptTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp.cleanup)
        self.root = Path(self.temp.name)
        self.state = self.root / 'state'; self.state.mkdir()
        self.base = self.root / 'baseline.xlsx'; self.base.write_bytes(b'baseline-test-only')
        self.registry = self.root / 'active.json'
        self.active = dict(status='ACTIVE', canonical_workbook=str(self.base), sha256=p.sha(self.base),
                           version='fixture', schema_version='fixture', updated_at='fixture')
        p.write_json(self.registry, self.active)
        self.cfg = dict(state=str(self.state), registry=str(self.registry))

    def batch(self, text=TEXT + STATS):
        source = self.root / ('source-' + str(len(list(self.root.glob('source-*')))) + '.txt')
        source.write_text(text, encoding='utf-8')
        out = p.prepare(self.cfg, [source], kind='transcript')
        path = Path(out['batch']); m = p.read_json(path)
        m['documents'][0]['review']['checked'] = True
        p.write_json(path, m)
        return path

    def test_absent_attributes_allowed(self):
        self.assertFalse(p.preflight_text(TEXT + STATS).attributes)

    def test_incomplete_attributes_rejected(self):
        with self.assertRaisesRegex(ValueError, 'ATTRIBUTE_SCHEMA_COUNT'):
            p.preflight_text(TEXT + '\nPLAYER_ATTRIBUTES\nSnapshot=2035-03-01\nAttribute_Schema=OUTFIELD\n傳球=12\n')

    def test_unreviewed_rejected(self):
        path = self.batch(); m = p.read_json(path); m['documents'][0]['review']['checked'] = False; p.write_json(path, m)
        with self.assertRaisesRegex(ValueError, 'NEEDS_REVIEW'):
            p.checked_batch(path)

    def test_uncertain_cell_rejected(self):
        path = self.batch(); m = p.read_json(path); m['documents'][0]['review']['unresolved'] = ['Goals: 8 or 9']; p.write_json(path, m)
        with self.assertRaisesRegex(ValueError, 'NEEDS_REVIEW'):
            p.checked_batch(path)

    def test_tampered_original_rejected(self):
        path = self.batch(); m = p.read_json(path); (path.parent / m['evidence'][0]['path']).write_bytes(b'changed')
        with self.assertRaisesRegex(ValueError, 'EVIDENCE_HASH_MISMATCH'):
            p.checked_batch(path)

    def test_path_escape_rejected(self):
        with self.assertRaisesRegex(ValueError, 'PATH_OUTSIDE_BATCH'):
            p.inside(self.state, '../outside.txt')

    def test_duplicate_documents_deduplicated(self):
        path = self.batch(); m = p.read_json(path); m['documents'] *= 2; p.write_json(path, m)
        self.assertEqual(len(p.checked_batch(path)[1]), 1)

    def test_original_image_preserved_and_duplicate_upload_deduplicated(self):
        from PIL import Image
        image = self.root / 'fixture.png'
        Image.new('RGB', (20, 30), 'white').save(image)
        out = p.prepare(self.cfg, [image, image])
        m = p.read_json(out['batch'])
        self.assertEqual(len(m['evidence']), 1)
        e = m['evidence'][0]
        self.assertEqual((e['width'], e['height']), (20, 30))
        self.assertEqual((Path(out['batch']).parent / e['path']).read_bytes(), image.read_bytes())

    def test_unaccounted_image_rejected(self):
        path = self.batch(); m = p.read_json(path)
        extra = path.parent / 'extra.txt'; extra.write_text('Another source', encoding='utf-8')
        m['evidence'].append(dict(path='extra.txt', sha256=p.sha(extra), kind='transcript'))
        p.write_json(path, m)
        with self.assertRaisesRegex(ValueError, 'UNACCOUNTED_SCREENSHOTS'):
            p.checked_batch(path)

    def test_unknown_section_rejected(self):
        with self.assertRaises(ValueError):
            p.preflight_text(TEXT + '\nAWARDS\nAward|Rank\nUnknown|1\n')

    def test_unmaterialized_fee_rejected(self):
        with self.assertRaisesRegex(ValueError, 'UNMATERIALIZED_SOURCE_FIELD:Fee'):
            p.preflight_text(TEXT + '\nLEAGUE_CAREER\nSeason|Club|Apps|Fee\n2034/35|Club|1|1000\n')

    def test_missing_competition_scope_rejected(self):
        with self.assertRaisesRegex(ValueError, 'COMPETITION_SCOPE_REQUIRED'):
            p.preflight_text(TEXT + STATS.replace('CLUB_TOTALS', 'CLUB_COMPETITION_STATS'))

    def test_unknown_national_team_level_rejected(self):
        with self.assertRaises(ValueError):
            p.preflight_text(TEXT + '\nNATIONAL_TEAM_COMPETITION_STATS\nSeason|National_Team_Raw|Team_Level|Context_Club|Apps_Raw\n2034/35|NULL|NULL|Club|3\n')

    def test_nonfinite_negative_fractional_counts_rejected(self):
        for token in ('NaN', 'inf', '-1', '1.5'):
            with self.subTest(token=token), self.assertRaisesRegex(ValueError, 'NUMBER_INVALID:Goals'):
                p.preflight_text(TEXT + STATS.replace('|8|0|', '|' + token + '|0|'))

    def test_source_null_and_zero_distinct(self):
        d = p.preflight_text(TEXT + STATS.replace('|8|0|', '|NULL|0|'))
        r = d.records['CLUB_TOTALS'][0]
        self.assertIsNone(r['Goals'].value); self.assertEqual(r['Assists'].value, '0')

    def test_appearance_semantics(self):
        self.assertEqual(parse_apps('10(3)'), (13, 10, 3))

    def test_stale_batch_does_not_change_active(self):
        path = self.batch(); m = p.read_json(path); m['baseline_sha256'] = 'old'; p.write_json(path, m)
        with self.assertRaisesRegex(ValueError, 'STALE_BATCH'):
            p.run_batch(self.cfg, path)
        self.assertEqual(p.read_json(self.registry), self.active)

    def fake_import(self, output, text, metadata, registry):
        before = p.active_registry(registry)
        output.write_bytes(Path(before['canonical_workbook']).read_bytes() + b'next')
        return dict(player_id='TEST', rows_added={'CLUB_TOTALS': 1}, selected_rows_updated=0, observations=0,
                    validation=dict(pass_=True, **{'pass': True}, unexpected_changes_count=0,
                                    baseline_sha256=before['sha256'], candidate_sha256=p.sha(output)))

    def fake_build(self, workbook, dest):
        dest.mkdir(); (dest / 'World_Master.xlsx').write_bytes(workbook.read_bytes()); (dest / 'archive.html').write_text('test')
        return dict(workbook_sha256=p.sha(workbook), sheets=1, rows=1)

    def fakes(self, **kwargs):
        stack = ExitStack()
        stack.enter_context(patch('production_importer.import_full_player_into_active', side_effect=kwargs.get('importer', self.fake_import)))
        stack.enter_context(patch.object(p, 'build_release', side_effect=kwargs.get('builder', self.fake_build)))
        stack.enter_context(patch.object(p, 'semantic_equal', return_value=False))
        stack.enter_context(patch.object(p, 'review_items', return_value=[]))
        return stack

    def test_build_failure_does_not_activate(self):
        def fail(*args): raise RuntimeError('simulated build failure')
        with self.fakes(builder=fail), self.assertRaisesRegex(RuntimeError, 'simulated'):
            p.run_batch(self.cfg, self.batch())
        self.assertEqual(p.read_json(self.registry), self.active)
        self.assertEqual(self.base.read_bytes(), b'baseline-test-only')

    def test_late_document_failure_rolls_back_entire_batch(self):
        path = self.batch(); m = p.read_json(path)
        second = path.parent / 'second.txt'; second.write_text((TEXT + STATS).replace('Fixture Player', 'Second Player'), encoding='utf-8')
        m['documents'].append(dict(m['documents'][0], path='second.txt')); p.write_json(path, m)
        calls = []
        def fail_second(*args):
            calls.append(1)
            if len(calls) == 2: raise RuntimeError('second document failed')
            return self.fake_import(*args)
        with self.fakes(importer=fail_second), self.assertRaisesRegex(RuntimeError, 'second document'):
            p.run_batch(self.cfg, path)
        self.assertEqual(len(calls), 2)
        self.assertEqual(p.read_json(self.registry), self.active)

    def test_activation_replay_and_rollback(self):
        path = self.batch()
        with self.fakes():
            result = p.run_batch(self.cfg, path)
            self.assertEqual(result['status'], 'ACTIVE')
            active = p.active_registry(self.registry)
            self.assertEqual(Path(active['canonical_workbook']).read_bytes(), b'baseline-test-onlynext')
            self.assertEqual(p.run_batch(self.cfg, path)['status'], 'NO_OP')
            p.rollback(self.cfg, active['sha256'])
        self.assertEqual(p.read_json(self.registry), self.active)

    def test_preview_does_not_activate(self):
        with self.fakes(): result = p.run_batch(self.cfg, self.batch(), activate=False)
        self.assertEqual(result['status'], 'VALIDATED')
        self.assertEqual(p.read_json(self.registry), self.active)

    def test_candidate_mutation_after_build_rejected(self):
        def mutate(workbook, dest):
            result = self.fake_build(workbook, dest); workbook.write_bytes(b'tampered'); return result
        with self.fakes(builder=mutate), self.assertRaisesRegex(ValueError, 'CANDIDATE_CHANGED'):
            p.run_batch(self.cfg, self.batch())
        self.assertEqual(p.read_json(self.registry), self.active)

    def test_release_copy_mutation_rejected(self):
        def mutate(workbook, dest):
            result = self.fake_build(workbook, dest); (dest / 'World_Master.xlsx').write_bytes(b'tampered'); return result
        with self.fakes(builder=mutate), self.assertRaisesRegex(ValueError, 'RELEASE_WORKBOOK_HASH_MISMATCH'):
            p.run_batch(self.cfg, self.batch())
        self.assertEqual(p.read_json(self.registry), self.active)

    def test_bad_gate_rejected(self):
        def bad(*args):
            r = self.fake_import(*args); r['validation']['candidate_sha256'] = 'bad'; return r
        with self.fakes(importer=bad), self.assertRaisesRegex(ValueError, 'CANDIDATE_GATE_FAILED'):
            p.run_batch(self.cfg, self.batch())
        self.assertEqual(p.read_json(self.registry), self.active)

    def test_active_changed_under_lock_rejected(self):
        def external_write(workbook, dest):
            result = self.fake_build(workbook, dest)
            p.write_json(self.registry, dict(self.active, updated_at='another writer'))
            return result
        with self.fakes(builder=external_write), self.assertRaisesRegex(ValueError, 'STALE_CANDIDATE_BASELINE'):
            p.run_batch(self.cfg, self.batch())
        self.assertEqual(p.read_json(self.registry)['canonical_workbook'], str(self.base))

    def test_wrong_rollback_version_rejected(self):
        with self.assertRaisesRegex(ValueError, 'ROLLBACK_BASELINE_CHANGED'):
            p.rollback(self.cfg, 'not-current')

    def test_parallel_writer_cannot_get_same_lock(self):
        with p.registry_lock(self.registry):
            with self.assertRaisesRegex(ValueError, 'IMPORT_ALREADY_RUNNING'):
                with p.registry_lock(self.registry):
                    self.fail('second writer acquired the lock')

    def test_secondary_report_failure_does_not_misreport_activation(self):
        original = p.write_json
        def fail_report(path, value):
            if Path(path).name == 'result.json': raise OSError('simulated report write error')
            return original(path, value)
        path = self.batch()
        with self.fakes(), patch.object(p, 'write_json', side_effect=fail_report):
            result = p.run_batch(self.cfg, path)
        self.assertEqual(result['status'], 'ACTIVE')
        self.assertIn('report_write_error', result)
        self.assertNotEqual(p.active_registry(self.registry)['sha256'], self.active['sha256'])

    def test_baseline_copy_mismatch_rejected(self):
        path = self.batch()
        with patch.object(p.shutil, 'copyfile', side_effect=lambda src, dst: Path(dst).write_bytes(b'wrong copy')):
            with self.assertRaisesRegex(ValueError, 'BASELINE_CHANGED_DURING_COPY'):
                p.run_batch(self.cfg, path)
        self.assertEqual(p.read_json(self.registry), self.active)


class AdoptionTests(unittest.TestCase):
    def test_equal_values_ignore_new_provenance(self):
        self.assertTrue(evidence_same(dict(Goals=2, Fact_ID='old'), dict(Goals=2, Fact_ID='new')))

    def test_latest_stats_keep_fact_id_and_new_observation(self):
        old = dict(Goals=2, Source_ID='old', Fact_ID='same-fact')
        incoming = dict(Goals=3, Source_ID='new')
        stamp(incoming, 'Player_Club_Season_Totals', ('p', 's', 'c'), 'abc', 2, FIELDS)
        chosen, rules = select_row(old, incoming, family='CLUB_TOTALS', old_snapshot='2035-01-01', snapshot='2035-02-01')
        self.assertEqual(chosen['Goals'], 3)
        self.assertEqual(chosen['Fact_ID'], 'same-fact')
        self.assertEqual(chosen['Source_Row_ID'], incoming['Source_Row_ID'])
        self.assertEqual(chosen['Statistical_Adoption_Status'], 'ADOPTED')

    def test_large_conflict_keeps_adopted_fact(self):
        old = dict(Goals=2, Source_ID='old', Fact_ID='same-fact')
        chosen, rules = select_row(old, dict(Goals=20, Source_ID='new'), family='CLUB_TOTALS', old_snapshot=None, snapshot='2035-02-01')
        self.assertEqual(chosen, old)
        self.assertEqual(rules['Goals'], 'LARGE_DIFF_ISOLATED_KEEP_EXISTING')


class ModernHonoursTests(unittest.TestCase):
    def fixture(self, season='PER-S-2034-35', adopted_status='ADOPTED'):
        from openpyxl import Workbook
        from barcelona_honours_sync_v1 import HEADERS, FIELDS
        w = Workbook(); w.remove(w.active)
        honours = w.create_sheet('Barcelona_Player_Season_Honours'); honours.append(HEADERS)
        honours.append(('P-OLD', 'PER-S-2034-35', 'C-0030', 2, 'A1', 1, 1, 0, 0, 1, 'curated', 'source', 'verified'))
        totals = w.create_sheet('Player_Club_Season_Totals')
        totals.append(('Player_ID', 'Season_ID', 'Club_ID', 'Apps', 'Statistical_Adoption_Status'))
        totals.append(('P-NEW', season, 'C-0030', 5, adopted_status))
        master = w.create_sheet('Barcelona_Season_Master'); master.append(('Season_ID',))
        career = w.create_sheet('Barcelona_Player_Career'); career.append(('Player_ID', *FIELDS, 'Major_Awards'))
        career.append(('P-OLD', 1, 1, 0, 0, 1, 'Latest curated awards'))
        career.append(('P-NEW', None, None, None, None, None, None))
        self.addCleanup(w.close)
        return w, FIELDS

    def test_existing_awards_preserved_and_known_season_copied(self):
        from modern_honours import sync
        w, fields = self.fixture()
        before = list(w['Barcelona_Player_Season_Honours'].values)[1]
        sync(w, fields, 'A1')
        self.assertEqual(list(w['Barcelona_Player_Season_Honours'].values)[1], before)
        self.assertEqual(w['Barcelona_Player_Career'].cell(2, 7).value, 'Latest curated awards')
        self.assertEqual(tuple(w['Barcelona_Player_Season_Honours'].values)[2][5:10], (1, 1, 0, 0, 1))
        self.assertEqual(sync(w, fields, 'A1')['season_honours_changes'], 0)

    def test_unknown_titles_do_not_become_zero(self):
        from modern_honours import sync
        w, fields = self.fixture(season='PER-S-2035-36')
        result = sync(w, fields, 'A1')
        self.assertTrue(result['pending'])
        self.assertEqual(tuple(w['Barcelona_Player_Season_Honours'].values)[2][5:10], (None,) * 5)

    def test_superseded_observation_adds_no_membership(self):
        from modern_honours import sync
        w, fields = self.fixture(adopted_status='SUPERSEDED')
        sync(w, fields, 'A1')
        self.assertEqual(w['Barcelona_Player_Season_Honours'].max_row, 2)

    def test_midseason_first_place_is_not_a_championship(self):
        from modern_honours import sync
        w, fields = self.fixture(season='PER-S-2035-36')
        ws = w['Barcelona_Season_Master']; ws.cell(1, 2, 'LaLiga_Position'); ws.cell(1, 3, 'Derivation_Status')
        ws.append(('PER-S-2035-36', 1, 'IN_PROGRESS'))
        sync(w, fields, 'A1')
        self.assertIsNone(w['Barcelona_Player_Season_Honours'].cell(3, 6).value)

    def test_superseded_rows_not_selected(self):
        self.assertFalse(adopted(('SUPERSEDED',), {'Statistical_Adoption_Status': 1}))
        self.assertTrue(adopted(('ADOPTED',), {'Statistical_Adoption_Status': 1}))

    def test_apps_group_not_split_by_selection(self):
        old = dict(Apps_Raw='10(3)', Apps=13, Starts=10, Sub_Appearances=3, Source_ID='a')
        new = dict(Apps_Raw='11(3)', Apps=14, Starts=11, Sub_Appearances=3, Source_ID='b')
        chosen, _ = select_row(old, new, family='CLUB_TOTALS', old_snapshot=None, snapshot='2035-02-01')
        self.assertEqual(chosen, old)


class DerivedMissingValueTests(unittest.TestCase):
    def test_new_career_unknown_assists_remain_null_and_rebuild_is_stable(self):
        from openpyxl import Workbook, load_workbook
        from derived_rebuild import rebuild_derived, SEASON_SAFE, CAREER_SAFE
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / 'synthetic.xlsx'
            w = Workbook(); w.remove(w.active)
            s = w.create_sheet('Player_Club_Season_Totals')
            s.append(('Season_ID', 'Player_ID', 'Club_ID', 'Apps', 'Starts', 'Goals', 'Assists', 'POTM', 'Rating', 'Statistical_Adoption_Status'))
            s.append(('PER-S-2034-35', 'TEST', 'C-0030', 5, 3, 2, None, 0, 7.25, 'ADOPTED'))
            w.create_sheet('Barcelona_Player_Season_Stats').append(('Season_ID', 'Player_ID', *SEASON_SAFE))
            w.create_sheet('Barcelona_Player_Career').append(('Player_ID', 'Player_Name', 'Derivation_Status', *CAREER_SAFE))
            s = w.create_sheet('Player_Dim'); s.append(('Player_ID', 'Canonical_Display_Name')); s.append(('TEST', 'Synthetic'))
            w.save(path); w.close()
            rebuild_derived(path, affected_players={'TEST'})
            w = load_workbook(path)
            rows = w['Barcelona_Player_Career'].iter_rows(values_only=True); cols = next(rows); row = dict(zip(cols, next(rows)))
            self.assertIsNone(row['Assists']); self.assertEqual(row['Goals'], 2); self.assertEqual(row['MOTM'], 0)
            w.close()
            self.assertEqual(rebuild_derived(path)['semantic_changes'], 0)


if __name__ == '__main__':
    unittest.main()
