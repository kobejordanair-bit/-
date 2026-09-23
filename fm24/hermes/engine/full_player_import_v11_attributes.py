"""Temporary-copy-only V1.1 attribute authority prototype; never changes ACTIVE registry/canonical."""
from __future__ import annotations
from pathlib import Path
from datetime import date
from openpyxl import load_workbook
from generic_importer import ImportFailure

OUTFIELD=('傳球','傳中','盯人','罰點球','技術','角球','界外球','盤帶','搶斷','任意球','射門','停球','頭球','遠射','才華','防守站位','工作投入','集中','決斷','領導力','侵略性','視野','團隊合作','無球跑動','意志力','勇敢','預判','鎮定','爆發力','彈跳','靈活','耐力','平衡','強壯','速度','體質')
GOALKEEPER=('出擊（傾向）','傳球','大腳開球','反應','攔截傳中','拳擊球（傾向）','手控球','手拋球','停球','一對一','意外性','指揮防守','制空範圍','才華','防守站位','工作投入','集中','決斷','領導力','侵略性','視野','團隊合作','無球跑動','意志力','勇敢','預判','鎮定','爆發力','彈跳','靈活','耐力','平衡','強壯','速度','體質','罰點球','技術','任意球')
HEADER=('Player_ID','Snapshot_Date','Attribute_Schema','Source_ID','Verification_Status')

def parse_attributes(lines, player_snapshot):
    """V1.1 section: Snapshot=YYYY-MM-DD; Attribute_Schema=...; exact Chinese attribute keys."""
    d={}
    for raw in lines:
        if '=' not in raw: raise ImportFailure('PLAYER_ATTRIBUTES_MALFORMED')
        k,v=(x.strip() for x in raw.split('=',1))
        if k in d: raise ImportFailure(f'PLAYER_ATTRIBUTES_DUPLICATE:{k}')
        d[k]=None if v in ('','NULL','-') else v
    schema=d.pop('Attribute_Schema',None); snapshot=d.pop('Snapshot',None)
    if schema not in ('OUTFIELD','GOALKEEPER'): raise ImportFailure('ATTRIBUTE_SCHEMA_INVALID')
    if snapshot != player_snapshot: raise ImportFailure('SNAPSHOT_DATE_MISMATCH')
    names=OUTFIELD if schema=='OUTFIELD' else GOALKEEPER
    if set(d)-set(names) or set(names)-set(d): raise ImportFailure('ATTRIBUTE_SCHEMA_COUNT_OR_NAME_INVALID')
    values={}
    for k,v in d.items():
        if v is None: raise ImportFailure('ATTRIBUTE_SNAPSHOT_INCOMPLETE')
        try: n=int(v)
        except ValueError: raise ImportFailure(f'ATTRIBUTE_VALUE_INVALID:{k}')
        if str(n)!=v or not 1<=n<=20: raise ImportFailure(f'ATTRIBUTE_OUT_OF_RANGE:{k}')
        values[k]=n
    return schema,snapshot,values

def ensure_attribute_sheets(path):
    wb=load_workbook(path)
    for schema,names in [('OUTFIELD',OUTFIELD),('GOALKEEPER',GOALKEEPER)]:
        sn='Player_Attr_Snap_'+('O' if schema=='OUTFIELD' else 'G')
        if sn not in wb.sheetnames: wb.create_sheet(sn).append(HEADER+names)
    if 'Player_Attribute_Changes' not in wb.sheetnames:
        wb.create_sheet('Player_Attribute_Changes').append(('Player_ID','From_Snapshot','To_Snapshot','Attribute_Name','Old_Value','New_Value','Delta'))
    wb.save(path)

def rebuild_changes(path):
    wb=load_workbook(path); out=wb['Player_Attribute_Changes']; out.delete_rows(2,out.max_row)
    added=0
    for schema,names in [('OUTFIELD',OUTFIELD),('GOALKEEPER',GOALKEEPER)]:
        ws=wb['Player_Attr_Snap_'+('O' if schema=='OUTFIELD' else 'G')]; rows=list(ws.iter_rows(min_row=2,values_only=True))
        for pid in sorted({r[0] for r in rows if r[0]}):
            hist=sorted((r for r in rows if r[0]==pid),key=lambda r:r[1])
            for a,b in zip(hist,hist[1:]):
                for i,n in enumerate(names,5):
                    if a[i]!=b[i]: out.append((pid,a[1],b[1],n,a[i],b[i],b[i]-a[i]));added+=1
    wb.save(path); return added
