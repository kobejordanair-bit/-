"""V1 stabilization wrapper. New module; legacy generic_importer remains untouched."""
from __future__ import annotations
from pathlib import Path
from shutil import copy2
from openpyxl import load_workbook
from generic_importer import (parse, resolve, source_id, headers, dictrow, refresh, parse_apps,
                              num, ImportFailure, TARGETS)


def _payload(rec, pid, cid, season, sid, status):
    apps, starts, subs = parse_apps(rec.values['Apps_Raw'])
    raw=lambda k: rec.raw[k].token if k in rec.raw else None
    d={'Season_ID':season,'Season_Display':rec.values['Season'],'Player_ID':pid,'Club_ID':cid,
       'Club_Raw':rec.values['Club'],'Apps_Raw':raw('Apps_Raw'),'Apps':apps,'Starts':starts,
       'Sub_Appearances':subs,'Goals':num(rec.values.get('Goals')),'Penalties_Raw':raw('Penalties'),
       'Assists':num(rec.values.get('Assists')),'POTM':num(rec.values.get('POTM')),
       'Yellow':num(rec.values.get('Yellow')),'Red':num(rec.values.get('Red')),
       'TacklesWon90':num(rec.values.get('TacklesWon90')),'PassPct_Raw':raw('PassPct'),
       'Dribbles90':num(rec.values.get('Dribbles90')),'ShotsOnTargetPct_Raw':raw('ShotsOnTargetPct'),
       'Fouls':num(rec.values.get('Fouls')),'Fouled':num(rec.values.get('Fouled')),
       'Rating':num(rec.values.get('Rating')),'Goals_Conceded':num(rec.values.get('Conceded')),
       'Clean_Sheets':num(rec.values.get('CleanSheets')),'Source_ID':sid,'Verification_Status':status}
    if rec.family=='CLUB_COMPETITION_STATS': d['Competition_Raw']=rec.values['Competition']
    return d

def _key(d, competition):
    ks=['Player_ID','Season_ID','Club_ID']+(['Competition_Raw'] if competition else [])
    return tuple(d[k] for k in ks)

def _same(row, h, payload):
    # Canonical equality compares every writable authority column, not formatting.
    return all(row[h[k]-1] == v for k,v in payload.items() if k in h)

def import_document_v1(base: Path, output: Path, text: str, metadata: dict):
    """Preflight every record; copy/save only after all conflicts are ruled out."""
    doc=parse(text)
    wb=load_workbook(base,data_only=False)
    resolved, coverage=resolve(doc,wb)
    sid,newsource=source_id(metadata,wb['Source_Index'])
    unsupported=[rec for rec,_,_,_ in resolved if 'Conceded90' in rec.raw and rec.values.get('Conceded90') is not None]
    if unsupported: raise ImportFailure('UNMATERIALIZED_SOURCE_FIELD: Conceded90 has no approved authority destination')
    planned=[]; noops=0
    for rec,pid,cid,season in resolved:
        sn=TARGETS[0] if rec.family=='CLUB_COMPETITION_STATS' else TARGETS[1]
        d=_payload(rec,pid,cid,season,sid,metadata.get('verification_status')); ws=wb[sn]; h=headers(ws)
        matches=[]
        for row in ws.iter_rows(min_row=2,values_only=True):
            if _key({k:row[h[k]-1] for k in h},rec.family=='CLUB_COMPETITION_STATS')==_key(d,rec.family=='CLUB_COMPETITION_STATS'):
                matches.append(row)
        if len(matches)>1: raise ImportFailure('BUSINESS_KEY_DUPLICATE_IN_WORKBOOK')
        if matches:
            old=matches[0]; old_sid=str(old[h['Source_ID']-1])
            if old_sid != sid: raise ImportFailure('SOURCE_CONFLICT_UNRESOLVED')
            if not _same(old,h,d): raise ImportFailure('SAME_SOURCE_RECORD_CONFLICT')
            noops+=1
        else: planned.append((sn,d))
    # Existing source only: idempotent rerun. A new source cannot be a pure no-op due conflicts above.
    if not planned:
        copy2(base,output)
        return {'source_id':sid,'new_source':False,'coverage':coverage,'records':len(resolved),
                'inserted_rows':0,'no_op_records':noops,'status':'ALREADY_IMPORTED_IDENTICAL'}
    copy2(base,output); wb=load_workbook(output,data_only=False)
    if newsource:
        ws=wb['Source_Index']; ws.append(dictrow(ws,{'Source_ID':sid,'File_Name':metadata['file_name'],
          'Local_Path':metadata['local_path'],'Domain':'FM World Import V1','Season_Context':doc.player.get('Snapshot'),
          'Source_Note':'execution provenance'})); refresh(ws)
    for sn,d in planned:
        ws=wb[sn]; ws.append(dictrow(ws,d)); refresh(ws)
    wb.save(output)
    return {'source_id':sid,'new_source':newsource,'coverage':coverage,'records':len(resolved),
            'inserted_rows':len(planned),'no_op_records':noops,'status':'IMPORTED'}
