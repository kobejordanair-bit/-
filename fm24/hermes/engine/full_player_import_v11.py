"""Full Player Import V1.1 formal production implementation.

The ACTIVE registry and its canonical workbook are read-only inputs.  V1.1 adds
attribute snapshots, explicit national-team identity/level, and deterministic
new-player allocation without performing a registry mutation.
"""
from __future__ import annotations

from collections import Counter, defaultdict
from dataclasses import dataclass
import hashlib
import json
from pathlib import Path
import re
from shutil import copyfile as copy2
import zipfile

from openpyxl import load_workbook
from screenshot_provenance import stamp, adopted

from barcelona_honours_sync_v1 import sync_workbook
from derived_rebuild import rebuild_derived
from full_player_import_v1_core import (
    Cell,
    PLAYER_FIELDS,
    SUMMARY_FIELDS,
    TABULAR,
    _key as _legacy_key,
    _payload,
    _reconcile,
    _same,
    _summary_payload,
    active_baseline,
    cell,
    date,
    number,
)
from full_player_import_v11_attributes import (
    GOALKEEPER,
    HEADER as ATTRIBUTE_HEADER,
    OUTFIELD,
    parse_attributes,
    rebuild_changes,
)
from generic_importer import ImportFailure, dictrow, headers, refresh, source_id, parse_apps
from identity_refresh import refresh_identity_map
from governed_identity import normalized_existing_candidates, verified_dobs
from import_policy import (ensure_observed_periods, evidence_same, select_row, observation,
                           append_observations, source_snapshot, dump, equivalent)

FORMAT = "FORMAT=FM_WORLD_IMPORT_V1_1"
SECTIONS = (
    "PLAYER",
    "PLAYER_ATTRIBUTES",
    "LEAGUE_CAREER",
    "LEAGUE_CAREER_TOTAL",
    "CLUB_COMPETITION_STATS",
    "CLUB_TOTALS",
    "NATIONAL_TEAM_COMPETITION_STATS",
)
CONTROLLED_NATION_ALIASES = {"英格兰": "N-0027"}
NATIONAL_FIELDS = set(TABULAR["NATIONAL_TEAM_COMPETITION_STATS"]) | {
    "National_Team_Raw",
    "Team_Level",
}


@dataclass(frozen=True)
class V11Document:
    player: dict
    player_raw: dict
    records: dict
    summary: dict
    attribute_schema: str
    attribute_snapshot: str
    attributes: dict


def parse_full_v11(text: str) -> V11Document:
    """Parse the complete V1.1 document and reject implicit national-team rows."""
    lines = [line.rstrip("\r") for line in text.splitlines()]
    if not lines or lines[0] != FORMAT:
        raise ImportFailure("FORMAT_VERSION_REQUIRED_OR_UNSUPPORTED")
    groups = defaultdict(list)
    current = None
    for raw in lines[1:]:
        stripped = raw.strip().lstrip("\ufeff")
        if stripped in SECTIONS:
            current = stripped
            continue
        if not stripped:
            continue
        if current is None:
            raise ImportFailure("CONTENT_OUTSIDE_SECTION")
        groups[current].append(raw)

    if "PLAYER" not in groups:
        raise ImportFailure("PLAYER_REQUIRED")

    player = {}
    player_raw = {}
    for line in groups["PLAYER"]:
        if "=" not in line:
            raise ImportFailure("PLAYER_MALFORMED")
        key, value = line.split("=", 1)
        key = key.strip()
        if key not in PLAYER_FIELDS:
            raise ImportFailure(f"PLAYER_FIELD_UNSUPPORTED:{key}")
        if key in player:
            raise ImportFailure(f"PLAYER_FIELD_DUPLICATE:{key}")
        parsed = cell(value)
        player[key], player_raw[key] = parsed.value, parsed.raw
    if player.get("Date_of_Birth"):
        date(player["Date_of_Birth"])
    if not player.get("Name"):
        raise ImportFailure("PLAYER_NAME_REQUIRED")
    if not player.get("Snapshot"):
        raise ImportFailure("PLAYER_SNAPSHOT_REQUIRED")
    date(player["Snapshot"])

    schema, attribute_snapshot, attributes = (None, None, {})
    if "PLAYER_ATTRIBUTES" in groups:
        schema, attribute_snapshot, attributes = parse_attributes(
            groups["PLAYER_ATTRIBUTES"], player["Snapshot"]
        )
    records = {}
    for section in TABULAR:
        block = groups.get(section, [])
        if not block:
            records[section] = []
            continue
        head = block[0].split("|")
        if section == "NATIONAL_TEAM_COMPETITION_STATS" and "Team_Level" not in head:
            raise ImportFailure("NATIONAL_TEAM_LEVEL_MISSING")
        allowed = (NATIONAL_FIELDS if section == "NATIONAL_TEAM_COMPETITION_STATS" else TABULAR[section]) | {"Goals_Conceded", "Clean_Sheets"}
        if len(set(head)) != len(head) or set(head) - allowed:
            raise ImportFailure(f"HEADER_UNSUPPORTED:{section}")
        if section == "NATIONAL_TEAM_COMPETITION_STATS":
            required = {"Season", "National_Team_Raw", "Team_Level", "Context_Club", "Apps_Raw"}
        elif section in {"LEAGUE_CAREER", "CLUB_COMPETITION_STATS", "CLUB_TOTALS"}:
            required = {"Season", "Club"}
        else:
            required = set()
        if not required <= set(head):
            raise ImportFailure(f"HEADER_INVALID:{section}")
        parsed_rows = []
        for line in block[1:]:
            values = line.split("|")
            if len(values) != len(head):
                raise ImportFailure(f"COLUMN_COUNT:{section}")
            row = {key: cell(value) for key, value in zip(head, values)}
            if any(row[key].value is None for key in required):
                raise ImportFailure(f"IDENTITY_NULL:{section}")
            if section == "NATIONAL_TEAM_COMPETITION_STATS" and row["Team_Level"].value not in {"Senior", "U21"}:
                raise ImportFailure(f"TEAM_LEVEL_INVALID:{row['Team_Level'].raw}")
            parsed_rows.append(row)
        records[section] = parsed_rows

    summary = {}
    for line in groups.get("LEAGUE_CAREER_TOTAL", []):
        if "=" not in line:
            raise ImportFailure("LEAGUE_CAREER_TOTAL_MALFORMED")
        key, value = line.split("=", 1)
        key = key.strip()
        if key not in SUMMARY_FIELDS:
            raise ImportFailure(f"LEAGUE_CAREER_TOTAL_FIELD_UNSUPPORTED:{key}")
        if key in summary:
            raise ImportFailure(f"LEAGUE_CAREER_TOTAL_FIELD_DUPLICATE:{key}")
        summary[key] = cell(value)
    return V11Document(player, player_raw, records, summary, schema, attribute_snapshot, attributes)


def _dimension_index(ws, id_field: str):
    h = headers(ws)
    result = defaultdict(set)
    for row in ws.iter_rows(min_row=2, values_only=True):
        for field in ("Canonical_Display_Name", "Alias_Name"):
            value = row[h[field] - 1]
            if value not in (None, ""):
                result[str(value)].add(str(row[h[id_field] - 1]))
    return result


def _one(index, token, label):
    from governed_identity import normalized_candidate_form
    normalized = normalized_candidate_form(str(token))
    values = set().union(*(ids for name,ids in index.items() if normalized_candidate_form(name)==normalized))
    if len(values) != 1:
        raise ImportFailure(f"{label}_{'NOT_FOUND' if not values else 'AMBIGUOUS'}:{token}")
    return next(iter(values))


def _nation_id(index, token, aliases_used):
    from governed_identity import normalized_candidate_form
    normalized = normalized_candidate_form(str(token))
    values = set().union(*(ids for name,ids in index.items() if normalized_candidate_form(name)==normalized))
    if len(values) == 1:
        return next(iter(values))
    if len(values) > 1:
        raise ImportFailure(f"NATION_AMBIGUOUS:{token}")
    controlled = CONTROLLED_NATION_ALIASES.get(token)
    if controlled is None:
        raise ImportFailure(f"NATION_NOT_FOUND:{token}")
    if controlled not in {value for ids in index.values() for value in ids}:
        raise ImportFailure(f"TEMPORARY_ALIAS_TARGET_NOT_FOUND:{token}:{controlled}")
    aliases_used[token] = controlled
    return controlled


def _period_index(wb):
    h = headers(wb["Period_Dim"])
    result = defaultdict(set)
    for row in wb["Period_Dim"].iter_rows(min_row=2, values_only=True):
        if row[h["Period_Type"] - 1] in {"Season", "Calendar Year"}:
            result[str(row[h["Source_Period_Display"] - 1])].add(str(row[h["Period_ID"] - 1]))
    return result


def _allocate_player(wb, player):
    tokens = list(dict.fromkeys([player['Name']] + ([player['Raw_Name']] if player.get('Raw_Name') else [])))
    resolved = [normalized_existing_candidates(wb, token) for token in tokens]
    if any(len(ids) > 1 for ids in resolved):
        raise ImportFailure('PLAYER_IDENTITY_AMBIGUOUS')
    ids = set().union(*resolved)
    if len(ids) > 1: raise ImportFailure('PLAYER_TOKENS_DIFFER')
    if ids:
        if any(not ids for ids in resolved): raise ImportFailure('PLAYER_TOKEN_UNRESOLVED')
        pid = next(iter(ids)); dobs = verified_dobs(wb, pid)
        if len(dobs) > 1 or player.get('Date_of_Birth') and dobs and dobs != {player['Date_of_Birth']}:
            raise ImportFailure('PLAYER_DOB_CONFLICT')
        return pid, False

    dob = player.get("Date_of_Birth")
    if not dob:
        raise ImportFailure("NEW_PLAYER_DOB_REQUIRED")
    collisions = set()
    snapshots = wb["Player_Profile_Snapshots"]
    h = headers(snapshots)
    dob_column = "Date_of_Birth" if "Date_of_Birth" in h else "DOB"
    for row in snapshots.iter_rows(min_row=2, values_only=True):
        if str(row[h[dob_column] - 1] or "") == dob:
            collisions.add(str(row[h["Player_ID"] - 1]))
    if collisions:
        raise ImportFailure(f"POSSIBLE_EXISTING_PLAYER_IDENTITY:{dob}:{','.join(sorted(collisions))}")
    numeric = []
    for row in wb["Player_Dim"].iter_rows(min_row=2, values_only=True):
        match = re.fullmatch(r"P-(\d+)", str(row[0] or ""))
        if match:
            numeric.append(int(match.group(1)))
    # Retired identities are tombstoned in the governed dimension notes.
    h = headers(wb['Player_Dim'])
    for row in wb['Player_Dim'].iter_rows(min_row=2, values_only=True):
        note = str(row[h['Source_Note']-1] or '')
        numeric.extend(int(x) for x in re.findall(r'"retired_player_id"\s*:\s*"P-(\d+)"', note))
    return f"P-{max(numeric, default=0) + 1:04d}", True


def _ensure_attribute_sheets(wb):
    for schema, names in (("OUTFIELD", OUTFIELD), ("GOALKEEPER", GOALKEEPER)):
        sheet_name = "Player_Attr_Snap_" + ("O" if schema == "OUTFIELD" else "G")
        expected = tuple(ATTRIBUTE_HEADER) + tuple(names)
        if sheet_name not in wb.sheetnames:
            wb.create_sheet(sheet_name).append(expected)
        elif tuple(cell.value for cell in wb[sheet_name][1]) != expected:
            raise ImportFailure(f"ATTRIBUTE_SHEET_SCHEMA_INVALID:{sheet_name}")
    change_header = (
        "Player_ID",
        "From_Snapshot",
        "To_Snapshot",
        "Attribute_Name",
        "Old_Value",
        "New_Value",
        "Delta",
    )
    if "Player_Attribute_Changes" not in wb.sheetnames:
        wb.create_sheet("Player_Attribute_Changes").append(change_header)
    elif tuple(cell.value for cell in wb["Player_Attribute_Changes"][1]) != change_header:
        raise ImportFailure("ATTRIBUTE_SHEET_SCHEMA_INVALID:Player_Attribute_Changes")


def _key(family, payload):
    from governed_identity import normalized_candidate_form
    d = dict(payload)
    competition = {'League':'LEAGUE', '联赛':'LEAGUE', 'Cup':'CUP', '杯赛':'CUP',
                   'Continental':'CONTINENTAL', '洲际级别':'CONTINENTAL'}
    for field in ('League_Raw', 'Competition_Raw', 'Context_Club'):
        if d.get(field) is not None:
            value = normalized_candidate_form(str(d[field]))
            d[field] = competition.get(value,value) if field=='Competition_Raw' else value
    if d.get('Team_Level') in ('成年隊（使用者確認預設）','成年隊（預設）'): d['Team_Level']='Senior'
    return _legacy_key(family, d)


def _existing_match(ws, family, payload):
    h = headers(ws)
    wanted = _key(family, payload)
    return [
        row
        for row in ws.iter_rows(min_row=2, values_only=True)
        if adopted(row, h) and _key(family, {name: row[column - 1] for name, column in h.items()}) == wanted
    ]


def _attribute_payload(doc, player_id, source_id_value, status):
    names = OUTFIELD if doc.attribute_schema == "OUTFIELD" else GOALKEEPER
    payload = {
        "Player_ID": player_id,
        "Snapshot_Date": doc.attribute_snapshot,
        "Attribute_Schema": doc.attribute_schema,
        "Source_ID": source_id_value,
        "Verification_Status": status,
    }
    payload.update({name: doc.attributes[name] for name in names})
    return payload


def _reconcile_national_summary(player, national_rows):
    aggregates = {}; classifications = {}
    for level, cap, goal in [('Senior','International_Caps','International_Goals'), ('U21','U21_Caps','U21_Goals')]:
        rows = [r for r in national_rows if r['Team_Level'].value == level]
        totals = {'Apps': None, 'Goals': None}
        if rows:
            apps = [parse_apps(r['Apps_Raw'].value)[0] for r in rows]
            goals = [r.get('Goals').value if r.get('Goals') else None for r in rows]
            totals = {'Apps':sum(apps) if all(x is not None for x in apps) else None,
                      'Goals':sum(int(x) for x in goals) if all(x is not None for x in goals) else None}
        declared = {'Apps':int(player[cap]) if player.get(cap) is not None else None,
                    'Goals':int(player[goal]) if player.get(goal) is not None else None}
        aggregates[level] = totals
        pairs = [(totals[k],declared[k]) for k in totals]
        if not rows or any(a is None or b is None for a,b in pairs): c = 'COVERAGE_UNKNOWN_NO_OVERWRITE'
        elif all(a == b for a,b in pairs): c = 'RECONCILED'
        elif all(a <= b for a,b in pairs): c = 'DIFFERENT_OR_INCOMPLETE_COVERAGE_NO_ACTION'
        else: c = 'ANOMALY_RETAIN_PROFILE_ISOLATE_COMPARISON'
        classifications[level] = c
    return {'classification': 'RECONCILED' if all(c == 'RECONCILED' for c in classifications.values()) else 'PRESERVED_BY_SCOPE_POLICY',
            **aggregates, 'classifications': classifications}


def import_full_v11(
    output: Path,
    text: str,
    metadata: dict,
    registry: Path = Path("/opt/data/FM24_World_Current.json"),
    *,
    input_workbook: Path | None = None,
):
    """Preflight completely, stage on a copy, then atomically replace ``output``."""
    canonical, registry_data = active_baseline(registry)
    base = canonical if input_workbook is None else Path(input_workbook)
    if input_workbook is not None:
        if base.resolve() == canonical.resolve():
            raise ImportFailure("INPUT_WORKBOOK_MUST_BE_TEMPORARY_COPY")
        if not base.is_file():
            raise ImportFailure(f"INPUT_WORKBOOK_NOT_FOUND:{base}")

    doc = parse_full_v11(text)
    wb = load_workbook(base, data_only=False)
    aliases_used = {}
    player_id, new_player = _allocate_player(wb, doc.player)
    clubs = _dimension_index(wb["Club_Dim"], "Club_ID")
    nations = _dimension_index(wb["Nation_Dim"], "National_Team_ID")
    period_additions = ensure_observed_periods(wb, [r['Season'].value for rows in doc.records.values() for r in rows])
    periods = _period_index(wb)
    current_club = _one(clubs, doc.player["Current_Club"], "CLUB") if doc.player.get("Current_Club") else None
    nationality = _nation_id(nations, doc.player["Nationality"], aliases_used) if doc.player.get("Nationality") else None
    sid, new_source = source_id(metadata, wb["Source_Index"])
    text_sha = hashlib.sha256(text.encode('utf-8')).hexdigest()
    source_headers = headers(wb['Source_Index'])
    for r in wb['Source_Index'].iter_rows(min_row=2, values_only=True):
        if f'TXT_SHA256={text_sha}' in str(r[source_headers['Source_Note']-1] or ''):
            sid, new_source = r[source_headers['Source_ID']-1], False
            break
    status = metadata.get("verification_status") or "SOURCE_TRANSCRIBED_V1_1"

    plans = [
        (
            "PLAYER",
            "Player_Profile_Snapshots",
            _payload("PLAYER", {}, player_id, current_club, None, nationality, sid, doc.player, status),
        )
    ]
    league_payloads = []
    for family, rows in doc.records.items():
        for row in rows:
            season = _one(periods, row["Season"].value, "PERIOD")
            if family == "NATIONAL_TEAM_COMPETITION_STATS":
                nation = _nation_id(nations, row["National_Team_Raw"].value, aliases_used)
                payload = _payload(family, row, player_id, None, season, nation, sid, doc.player, status)
                payload["Country_Raw"] = row["National_Team_Raw"].raw
                payload["Team_Level"] = row["Team_Level"].value
                sheet = "Player_National_Team_Stats"
            else:
                club = _one(clubs, row["Club"].value, "CLUB")
                payload = _payload(family, row, player_id, club, season, nationality, sid, doc.player, status)
                sheet = {
                    "LEAGUE_CAREER": "Player_League_Career",
                    "CLUB_COMPETITION_STATS": "Player_Club_Competition_Stats",
                    "CLUB_TOTALS": "Player_Club_Season_Totals",
                }[family]
            for stat in ('Goals_Conceded', 'Clean_Sheets'):
                if stat in row: payload[stat] = number(row[stat].value, stat)
            plans.append((family, sheet, payload))
            if family == "LEAGUE_CAREER":
                league_payloads.append(payload)
    summary_payload = _summary_payload(doc.player, player_id, sid, status, doc.summary)
    if summary_payload:
        plans.append(("LEAGUE_CAREER_TOTAL", "Player_Career_Summaries", summary_payload))

    inserts = []
    updates = []
    observations = []
    noops = defaultdict(int)
    seen_plans = set()
    for ordinal, (family, sheet, payload) in enumerate(plans, 1):
        stamp(payload, sheet, _key(family, payload), text_sha, ordinal, headers(wb[sheet]))
        key = (family, _key(family, payload))
        if key in seen_plans: raise ImportFailure('DUPLICATE_INPUT_BUSINESS_KEY:' + family)
        seen_plans.add(key)
        matches = _existing_match(wb[sheet], family, payload)
        if len(matches) > 1:
            raise ImportFailure(f'BUSINESS_KEY_DUPLICATE_IN_WORKBOOK:{family}')
        if matches:
            h = headers(wb[sheet]); old = {k: matches[0][v-1] for k,v in h.items()}
            if evidence_same(old, payload):
                noops[family] += 1
            else:
                chosen, rules = select_row(old, payload, family=family,
                    old_snapshot=source_snapshot(wb, old.get('Source_ID')), snapshot=doc.player['Snapshot'])
                observations.append(observation(family, sheet, old, payload, chosen, rules, doc.player['Snapshot']))
                if chosen != old: updates.append((family, sheet, chosen))
                else: noops[family] += 1
        else:
            inserts.append((family, sheet, payload))

    if doc.attributes:
        attribute_sheet = "Player_Attr_Snap_" + ("O" if doc.attribute_schema == "OUTFIELD" else "G")
        attribute_payload = _attribute_payload(doc, player_id, sid, status)
        if attribute_sheet in wb.sheetnames:
            ah = headers(wb[attribute_sheet])
            matches = [
                row
                for row in wb[attribute_sheet].iter_rows(min_row=2, values_only=True)
                if row[ah["Player_ID"] - 1] == player_id
                and str(row[ah["Snapshot_Date"] - 1]) == doc.attribute_snapshot
            ]
            if matches:
                if len(matches) > 1:
                    raise ImportFailure("BUSINESS_KEY_DUPLICATE_IN_WORKBOOK:PLAYER_ATTRIBUTES")
                if not evidence_same({k:matches[0][v-1] for k,v in ah.items()}, attribute_payload):
                    raise ImportFailure("SAME_SOURCE_RECORD_CONFLICT:PLAYER_ATTRIBUTES")
                noops["PLAYER_ATTRIBUTES"] += 1
            else:
                inserts.append(("PLAYER_ATTRIBUTES", attribute_sheet, attribute_payload))
        else:
            inserts.append(("PLAYER_ATTRIBUTES", attribute_sheet, attribute_payload))

    reconciliation = _reconcile(league_payloads, summary_payload) if summary_payload else None
    national_reconciliation = _reconcile_national_summary(
        doc.player, doc.records.get("NATIONAL_TEAM_COMPETITION_STATS", [])
    )
    output = Path(output)
    if output.resolve() in {canonical.resolve(), base.resolve()}:
        raise ImportFailure('OUTPUT_MUST_BE_ISOLATED_CANDIDATE')
    if not (inserts or updates or observations or period_additions or new_player):
        # A verified no-op needs no workbook serialization or derived rebuild.
        # Byte equality is the gate here; this does not certify old baseline issues.
        baseline_sha = hashlib.sha256(base.read_bytes()).hexdigest()
        copy2(base, output)
        candidate_sha = hashlib.sha256(output.read_bytes()).hexdigest()
        if candidate_sha != baseline_sha:
            raise ImportFailure('NO_OP_COPY_HASH_MISMATCH')
        wb.close()
        return {'baseline': str(base), 'canonical': str(canonical), 'canonical_sha256': registry_data['sha256'],
                'player_id': player_id, 'new_player': False, 'observations': 0, 'selected_rows_updated': 0,
                'periods_added': 0, 'controlled_aliases_used': aliases_used, 'source_id': sid,
                'rows_added': {family: 0 for family in ('PLAYER', 'PLAYER_ATTRIBUTES', 'LEAGUE_CAREER',
                    'LEAGUE_CAREER_TOTAL', 'CLUB_COMPETITION_STATS', 'CLUB_TOTALS', 'NATIONAL_TEAM_COMPETITION_STATS')},
                'noops': dict(noops), 'attribute_changes': 0, 'validation_pass': True,
                'validation': {'pass': True, 'mode': 'UNCHANGED_BYTES', 'baseline_sha256': baseline_sha,
                    'candidate_sha256': candidate_sha, 'unexpected_changes': [], 'unexpected_changes_count': 0}}
    stage = Path(str(output) + ".staging.xlsx")
    if stage.exists():
        stage.unlink()
    try:
        copy2(base, stage)
        staged = load_workbook(stage, data_only=False)
        _ensure_attribute_sheets(staged)
        if new_player:
            player_dim = staged["Player_Dim"]
            player_dim.append(
                dictrow(
                    player_dim,
                    {
                        "Player_ID": player_id,
                        "Canonical_Display_Name": doc.player["Name"],
                        "Alias_Name": doc.player.get("Raw_Name") or doc.player["Name"],
                        "Alias_Type": "FM_WORLD_IMPORT_V1_1 onboarding",
                        "Verification_Status": status,
                        "Source_Note": "FM_WORLD_IMPORT_V1_1 production import",
                    },
                )
            )
            refresh(player_dim)
        if (inserts or observations or updates) and new_source:
            source = staged["Source_Index"]
            source.append(
                dictrow(
                    source,
                    {
                        "Source_ID": sid,
                        "File_Name": metadata["file_name"],
                        "Local_Path": metadata["local_path"],
                        "Domain": "FM World Import V1.1 Full Player",
                        "Season_Context": doc.player["Snapshot"],
                        "Source_Note": f"production execution provenance; TXT_SHA256={text_sha}",
                    },
                )
            )
            refresh(source)
        for payload in period_additions:
            staged['Period_Dim'].append(dictrow(staged['Period_Dim'], payload))
        if period_additions: refresh(staged['Period_Dim'])
        for family, sheet, payload in updates:
            target = staged[sheet]; h = headers(target)
            for r in range(2, target.max_row+1):
                old = {k:target.cell(r,c).value for k,c in h.items()}
                if old.get("Statistical_Adoption_Status", "ADOPTED") == "ADOPTED" and _key(family, old) == _key(family, payload):
                    for k,v in payload.items(): target.cell(r,h[k]).value = v
                    break
        append_observations(staged, observations)
        for _, sheet, payload in inserts:
            target = staged[sheet]
            target.append(dictrow(target, payload))
            refresh(target)
        staged.save(stage)

        identity = refresh_identity_map(stage, reconcile_row_provenance=True)
        affected = {payload['Player_ID'] for _, sheet, payload in inserts + updates
                    if sheet == 'Player_Club_Season_Totals'}
        derived = rebuild_derived(stage, affected_players=affected)
        # Every successful Full Player import completes the common derived phase.
        # It is idempotent when no Barcelona membership exists or nothing is missing.
        honours = sync_workbook(stage)
        attribute_changes = rebuild_changes(stage)
        with zipfile.ZipFile(stage) as archive:
            bad_member = archive.testzip()
        if bad_member is not None:
            raise ImportFailure(f"XLSX_INTEGRITY_FAILED:{bad_member}")
        load_workbook(stage, read_only=True).close()
        validation = validate_v11_production(stage, base)
        if not validation["pass"]:
            raise ImportFailure(f"V11_FORMAL_VALIDATION_FAILED:{validation}")
        stage.replace(output)
    except Exception:
        if stage.exists():
            stage.unlink()
        raise

    families = (
        "PLAYER",
        "PLAYER_ATTRIBUTES",
        "LEAGUE_CAREER",
        "LEAGUE_CAREER_TOTAL",
        "CLUB_COMPETITION_STATS",
        "CLUB_TOTALS",
        "NATIONAL_TEAM_COMPETITION_STATS",
    )
    return {
        "baseline": str(base),
        "canonical": str(canonical),
        "canonical_sha256": registry_data["sha256"],
        "player_id": player_id,
        "new_player": new_player,
        "observations": len(observations),
        "selected_rows_updated": len(updates),
        "periods_added": len(period_additions),
        "controlled_aliases_used": aliases_used,
        "source_id": sid,
        "rows_added": {family: sum(item[0] == family for item in inserts) for family in families},
        "noops": dict(noops),
        "attribute_changes": attribute_changes,
        "identity": identity,
        "derived": derived,
        "honours": honours,
        "reconciliation": reconciliation,
        "national_reconciliation": national_reconciliation,
        "validation": validation,
        "validation_pass": True,
    }



def validate_v11_production(candidate: Path, baseline: Path) -> dict:
    """Run the formal, validator-owned V1.1 populated-candidate gate."""
    from global_validation import validate_v11
    return validate_v11(Path(baseline), Path(candidate), mode="populated")

def canonical_sha256(registry: Path = Path("/opt/data/FM24_World_Current.json")):
    data = json.loads(Path(registry).read_text(encoding="utf-8"))
    return hashlib.sha256(Path(data["canonical_workbook"]).read_bytes()).hexdigest()
