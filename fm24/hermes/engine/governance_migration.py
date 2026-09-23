"""Controlled governance operations, restricted to workbook copies."""
from __future__ import annotations
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries
from generic_importer import headers, dictrow, refresh, ImportFailure

PLAYER_SEED={'杨·阿尔诺':'P-0166','赖宇豪':'P-0288','让·阿史':'P-0110'}
CLUB_SEED={'托特纳姆':'C-0036','巴塞罗那二队':'C-0180','帕尔梅拉斯':'C-0083'}

def _index(ws,idfield):
 h=headers(ws); o={}
 for r in ws.iter_rows(min_row=2,values_only=True):
  for field in ('Canonical_Display_Name','Alias_Name'):
   if r[h[field]-1] not in (None,''): o.setdefault(str(r[h[field]-1]),set()).add(str(r[h[idfield]-1]))
 return o

def apply_aliases(path:Path, players=PLAYER_SEED, clubs=CLUB_SEED):
 wb=load_workbook(path); result={}
 for sheet,idfield,proposals in [('Player_Dim','Player_ID',players),('Club_Dim','Club_ID',clubs)]:
  ws=wb[sheet]; h=headers(ws); existing=_index(ws,idfield); added=[]
  for token,ident in proposals.items():
   if token in existing:
    if existing[token]!={ident}: raise ImportFailure(f'ALIAS_CONFLICT:{sheet}:{token}')
    continue
   source=next(r for r in ws.iter_rows(min_row=2,values_only=True) if str(r[h[idfield]-1])==ident)
   row=list(source); row[h['Alias_Name']-1]=token; row[h['Alias_Type']-1]='controlled authority-evidenced alias (temporary)'
   ws.append(row); added.append({'raw_token':token,'entity_id':ident,'canonical':source[h['Canonical_Display_Name']-1]})
  refresh(ws); result[sheet]=added
 wb.save(path); return result

def apply_bayern_c0188_do_not_use(path:Path):
 """Retain legacy C-0188 but mark it unavailable for new resolver references."""
 wb=load_workbook(path); ws=wb['Club_Dim']; h=headers(ws); changed=[]
 for rn,r in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
  if str(r[h['Club_ID']-1])=='C-0188':
   for field,value in [('Verification_Status','DEPRECATED — do not use for new references'),('Source_Note','Legacy duplicate of C-0039 retained for provenance; resolver must reject new references.')]:
    if ws.cell(rn,h[field]).value!=value:
     ws.cell(rn,h[field]).value=value; changed.append(f'{ws.title}!{rn}:{field}')
 if len(changed)!=2: raise ImportFailure(f'BAYERN_POLICY_PATTERN_UNSAFE:{len(changed)}')
 wb.save(path); return {'pass':True,'retained_id':'C-0188','active_id':'C-0039','changed_cells':changed}

def repair_jean_apps(path:Path):
 wb=load_workbook(path); ws=wb['Player_Club_Competition_Stats']; h=headers(ws); changes=[]
 for rn,r in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
  if str(r[h['Player_ID']-1])=='P-0110' and str(r[h['Apps_Raw']-1])=='0' and r[h['Apps']-1]==0 and r[h['Starts']-1] is None and r[h['Sub_Appearances']-1] is None:
   for f in ('Starts','Sub_Appearances'): ws.cell(rn,h[f]).value=0; changes.append(f'{ws.title}!{rn}:{f}')
 if len(changes)!=12: raise ImportFailure(f'JEAN_REPAIR_PATTERN_UNSAFE:{len(changes)}')
 wb.save(path); return {'pass':True,'changed_cells':len(changes),'other_changes':0,'cells':changes}

def validate_duplicate_table_consistency(path:Path):
 wb=load_workbook(path); violations=[]
 for ws in wb.worksheets:
  tables=list(ws.tables.values())
  for t in tables:
   if range_boundaries(t.ref)[3] != ws.max_row: violations.append(f'{ws.title}:{t.name}:ref_not_full_data_range:{t.ref}:{ws.max_row}')
  for i,t in enumerate(tables):
   for u in tables[i+1:]:
    a=range_boundaries(t.ref); b=range_boundaries(u.ref)
    overlap=not(a[2]<b[0] or b[2]<a[0] or a[3]<b[1] or b[3]<a[1])
    if overlap and t.ref!=u.ref: violations.append(f'{ws.title}:{t.name}/{u.name}:overlap_refs_differ:{t.ref}!={u.ref}')
 return {'pass':not violations,'violations':violations}
