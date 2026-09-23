#!/usr/bin/env python3
"""Read-only Global Validation gate for a scoped World Master candidate.

The candidate is never mutated: Identity Refresh and Derived Sync V1 idempotency are
executed only on disposable copies. Any semantic delta outside the declared scope
is a hard failure.
"""
from __future__ import annotations
import argparse, hashlib, json, shutil, sys, tempfile, zipfile
from pathlib import Path
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
from generic_importer import TARGETS, app_violations, validate as generic_validate, headers
from governance_migration import validate_duplicate_table_consistency
from identity_refresh import refresh_identity_map
from derived_rebuild import rebuild_derived, formula_errors
from barcelona_honours_sync_v1 import (SHEET as HONOURS_SHEET, HEADERS as HONOURS_HEADERS,
                                       FIELDS as HONOURS_FIELDS, membership_rows)

RAW_SHEETS = [
    "Barcelona_History", "Barcelona_UCL_Journey", "National_Tournaments", "UEFA_Super_Cup",
    "Barcelona_Club_World_Cup", "FIFA_FIFPro_World_XI", "Ballon_dOr", "Kopa_Trophy",
    "Youth_Awards", "UCL_Golden_Boot", "Pichichi_Award", "Barcelona_LaLiga_Best_XI",
    "Player_Club_Season_Totals", "Record_Identity_Map", "Source_Index",
]

def fingerprint(ws):
    h=hashlib.sha256()
    for row in ws.iter_rows():
        for c in row: h.update(repr(c.value).encode('utf-8')); h.update(b'\x1e')
        h.update(b'\x1f')
    return h.hexdigest()

def semantic_rows(path):
    wb=load_workbook(path,data_only=False)
    out={}
    for ws in wb.worksheets:
        out[ws.title]=[[c.value for c in row] for row in ws.iter_rows()]
    return out

def cell_deltas(base, candidate):
    out=[]
    for sn in base:
        a,b=base[sn],candidate[sn]
        mr=max(len(a),len(b)); mc=max(max((len(x) for x in a),default=0),max((len(x) for x in b),default=0))
        for r in range(mr):
            for c in range(mc):
                av=a[r][c] if r<len(a) and c<len(a[r]) else None
                bv=b[r][c] if r<len(b) and c<len(b[r]) else None
                if av!=bv: out.append((sn,r+1,c+1,av,bv))
    return out

def natural_dupes(ws, fields):
    h=headers(ws); seen=set(); n=0
    for r in ws.iter_rows(min_row=2,values_only=True):
        key=tuple(r[h[f]-1] for f in fields)
        if any(v is None for v in key): continue
        if key in seen: n+=1
        seen.add(key)
    return n

def fk_errors(path):
    wb=load_workbook(path,data_only=False); errors=[]
    pd={r[0] for r in wb['Player_Dim'].iter_rows(min_row=2,values_only=True)}
    cd={r[0] for r in wb['Club_Dim'].iter_rows(min_row=2,values_only=True)}
    per={r[0] for r in wb['Period_Dim'].iter_rows(min_row=2,values_only=True)}
    si={r[0] for r in wb['Source_Index'].iter_rows(min_row=2,values_only=True)}
    for sn in TARGETS:
        ws=wb[sn]; h=headers(ws)
        for i,r in enumerate(ws.iter_rows(min_row=2,values_only=True),2):
            for fld, pool in [('Player_ID',pd),('Club_ID',cd),('Season_ID',per),('Source_ID',si)]:
                if r[h[fld]-1] not in pool: errors.append(f'{sn}!{i}:{fld}')
    return errors

def validate_barcelona_honours_v1(path):
    """V1 invariants: A1 coverage, unique keys, valid identities and Career sums.

    CWC is deliberately not represented in this auto-sync sheet or reconciliation.
    """
    wb=load_workbook(path,data_only=False)
    if HONOURS_SHEET not in wb.sheetnames:
        return {'pass':False,'error':'MISSING_BARCELONA_PLAYER_SEASON_HONOURS'}
    ws=wb[HONOURS_SHEET]; h=headers(ws)
    if tuple(c.value for c in ws[1]) != HONOURS_HEADERS:
        return {'pass':False,'error':'HONOURS_SCHEMA_MISMATCH'}
    players={r[0] for r in wb['Player_Dim'].iter_rows(min_row=2,values_only=True)}
    seasons={r[0] for r in wb['Period_Dim'].iter_rows(min_row=2,values_only=True)}
    keys=[]; blank=unresolved=invalid=0; sums={}
    for r in ws.iter_rows(min_row=2,values_only=True):
        pid,season=r[h['Player_ID']-1],r[h['Season_ID']-1]
        if pid in (None,''): blank+=1
        elif pid not in players: unresolved+=1
        if season not in seasons: invalid+=1
        keys.append((pid,season))
        if pid not in (None,''):
            bucket=sums.setdefault(pid,{f:0 for f in HONOURS_FIELDS})
            for f in HONOURS_FIELDS: bucket[f]+=r[h[f]-1] or 0
    coverage={(r['Player_ID'],r['Season_ID']) for r in membership_rows(wb)}
    career=wb['Barcelona_Player_Career']; ch=headers(career); mismatches=[]
    for r in career.iter_rows(min_row=2,values_only=True):
        pid=r[ch['Player_ID']-1]
        for f in HONOURS_FIELDS:
            if r[ch[f]-1] != sums.get(pid,{}).get(f,0): mismatches.append(f'{pid}:{f}')
    return {'pass':not blank and not unresolved and not invalid and len(keys)==len(set(keys)) and coverage==set(keys) and not mismatches and 'Club_World_Cup_Team_Championships_While_At_Barcelona' not in h,
            'natural_key_duplicates':len(keys)-len(set(keys)),'blank_player_id_rows':blank,
            'unresolved_player_id_rows':unresolved,'invalid_season_id_rows':invalid,
            'coverage_missing':len(coverage-set(keys)),'coverage_extra':len(set(keys)-coverage),
            'career_reconciliation_mismatches':mismatches,'cwc_excluded': 'Club_World_Cup_Team_Championships_While_At_Barcelona' not in h}

def allowed_delta(delta):
    sn,r,c,old,new=delta
    if sn=='Barcelona_Player_Career' and r>=2:
        # P-0167 row only, approved honours plus derivation metadata.
        return r==16 and c in set(range(9,20))
    if sn=='Player_Career_Summaries':
        return r in (29,30,31) and c in range(1,14) and old is None
    return False

def main():
    ap=argparse.ArgumentParser(); ap.add_argument('--baseline',required=True,type=Path); ap.add_argument('--candidate',required=True,type=Path); ap.add_argument('--report',required=True,type=Path); args=ap.parse_args()
    before=semantic_rows(args.baseline); after=semantic_rows(args.candidate)
    deltas=cell_deltas(before,after)
    raw_base={s:fingerprint(load_workbook(args.baseline,data_only=False)[s]) for s in RAW_SHEETS}
    raw_out={s:fingerprint(load_workbook(args.candidate,data_only=False)[s]) for s in RAW_SHEETS}
    generic=generic_validate(args.candidate)
    wb=load_workbook(args.candidate,data_only=False)
    dups={sn:natural_dupes(wb[sn],fields) for sn,fields in [(TARGETS[0],['Player_ID','Club_ID','Competition_Raw','Season_ID']),(TARGETS[1],['Player_ID','Club_ID','Season_ID'])]}
    with tempfile.TemporaryDirectory(dir=HERE,prefix='.global_validation_') as td:
        probe=Path(td)/args.candidate.name; shutil.copy2(args.candidate,probe)
        identity1=refresh_identity_map(probe); identity2=refresh_identity_map(probe); derived=rebuild_derived(probe)
        from barcelona_honours_sync_v1 import sync_workbook
        honours_first=sync_workbook(probe); honours_second=sync_workbook(probe)
    # Candidate-specific exact checks.
    career=wb['Barcelona_Player_Career']; ch=headers(career); target=[r for r in range(2,career.max_row+1) if career.cell(r,ch['Player_ID']).value=='P-0167']
    expected_honors={'LaLiga_Team_Championships_While_At_Barcelona':1,'UCL_Team_Championships_While_At_Barcelona':2,'Copa_Team_Championships_While_At_Barcelona':2,'Supercopa_Team_Championships_While_At_Barcelona':2,'UEFA_SuperCup_Team_Championships_While_At_Barcelona':2,'Club_World_Cup_Team_Championships_While_At_Barcelona':1,'FIFPro_Count':0,'Ballon_dOr_Top3_Count':0,'Major_Awards':'2033/34 LaLiga Best XI（入選）'}
    honor_ok=len(target)==1 and all(career.cell(target[0],ch[f]).value==v for f,v in expected_honors.items())
    sws=wb['Player_Career_Summaries']; sh=headers(sws); summary=[r for r in sws.iter_rows(min_row=2,values_only=True) if r[sh['Player_ID']-1]=='P-0167']
    summary_ok=len(summary)==3
    honours_validation=validate_barcelona_honours_v1(args.candidate)
    report={'baseline':str(args.baseline),'candidate':str(args.candidate),'generic_validation':generic,'natural_key_duplicates':dups,'foreign_key_errors':fk_errors(args.candidate),'formula_errors':formula_errors(args.candidate),'table_consistency':validate_duplicate_table_consistency(args.candidate),'apps_derivation_violations':app_violations(args.candidate),'identity_refresh_first':identity1,'identity_refresh_second':identity2,'derived_sync_v1_probe':derived,'honours_sync_first_probe':honours_first,'honours_sync_second_probe':honours_second,'barcelona_honours_v1':honours_validation,'raw_fingerprints_unchanged':raw_base==raw_out,'semantic_delta_count':len(deltas),'unexpected_semantic_deltas':[d for d in deltas if not allowed_delta(d)],'archie_honors_exact':honor_ok,'archie_summary_rows':len(summary),'archie_summary_count_exact':summary_ok}
    report['pass']=all(generic.values()) and all(v==0 for v in dups.values()) and not report['foreign_key_errors'] and not report['formula_errors'] and report['table_consistency']['pass'] and not report['apps_derivation_violations'] and identity1['added_map_entries']==0 and identity2['added_map_entries']==0 and derived['semantic_changes']==0 and honours_validation['pass'] and honours_first['season_honours_changes']==0 and honours_first['first_run_career_changes']==0 and honours_second['season_honours_changes']==0 and honours_second['first_run_career_changes']==0 and report['raw_fingerprints_unchanged'] and not report['unexpected_semantic_deltas'] and honor_ok and summary_ok
    args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8')
    print(json.dumps({'pass':report['pass'],'semantic_delta_count':len(deltas),'unexpected_semantic_deltas':len(report['unexpected_semantic_deltas']),'archie_summary_rows':len(summary),'generic_validation':generic},ensure_ascii=False))
    if not report['pass']: raise SystemExit(1)
# V1.1 validation additions are defined below; CLI dispatch is at end of file.

# --- Full Player Import V1.1 formal validation --------------------------------
from datetime import date as _date
import re
from full_player_import_v11_attributes_prototype import HEADER as _ATTR_HEADER, OUTFIELD as _OUTFIELD, GOALKEEPER as _GOALKEEPER

_V11_SHEETS = {
    "Player_Attr_Snap_O": tuple(_ATTR_HEADER) + tuple(_OUTFIELD),
    "Player_Attr_Snap_G": tuple(_ATTR_HEADER) + tuple(_GOALKEEPER),
    "Player_Attribute_Changes": ("Player_ID", "From_Snapshot", "To_Snapshot", "Attribute_Name", "Old_Value", "New_Value", "Delta"),
}


def _v11_date(value):
    if isinstance(value, _date): return value
    if not isinstance(value, str): raise ValueError(value)
    return _date.fromisoformat(value)


def _v11_rows(ws):
    h = headers(ws)
    return [{name: row[col - 1] for name, col in h.items()} for row in ws.iter_rows(min_row=2, values_only=True) if any(v not in (None, "") for v in row)]


def _v11_player_ids(wb):
    h = headers(wb["Player_Dim"])
    return [str(r[h["Player_ID"] - 1]) for r in wb["Player_Dim"].iter_rows(min_row=2, values_only=True) if r[h["Player_ID"] - 1] not in (None, "")]


def validate_v11_attribute_sheets(path: Path) -> dict:
    """Validate exact V1.1 snapshot authority schemas and derived delta semantics."""
    wb = load_workbook(path, data_only=False)
    errors=[]; player_ids=set(_v11_player_ids(wb)); snapshots={}
    for schema, sheet, names in (("OUTFIELD", "Player_Attr_Snap_O", _OUTFIELD), ("GOALKEEPER", "Player_Attr_Snap_G", _GOALKEEPER)):
        if sheet not in wb.sheetnames or tuple(c.value for c in wb[sheet][1]) != _V11_SHEETS[sheet]:
            errors.append(f"ATTRIBUTE_SHEET_SCHEMA_INVALID:{sheet}"); continue
        keys=set()
        for row in _v11_rows(wb[sheet]):
            key=(str(row.get("Player_ID")), str(row.get("Snapshot_Date")), str(row.get("Attribute_Schema")))
            if row.get("Player_ID") not in player_ids: errors.append("ATTRIBUTE_PLAYER_ID_INVALID")
            try: _v11_date(row.get("Snapshot_Date"))
            except (TypeError, ValueError): errors.append("ATTRIBUTE_SNAPSHOT_DATE_INVALID")
            if row.get("Attribute_Schema") != schema: errors.append("ATTRIBUTE_SCHEMA_INVALID")
            if key in keys: errors.append("ATTRIBUTE_NATURAL_KEY_DUPLICATE")
            keys.add(key)
            for name in names:
                value=row.get(name)
                if value in (None, ""): errors.append("ATTRIBUTE_SNAPSHOT_INCOMPLETE"); break
                if type(value) is not int: errors.append("ATTRIBUTE_VALUE_INVALID"); break
                if not 1 <= value <= 20: errors.append("ATTRIBUTE_OUT_OF_RANGE"); break
            snapshots[(str(row.get("Player_ID")), schema, str(row.get("Snapshot_Date")))] = row
    sheet="Player_Attribute_Changes"
    if sheet not in wb.sheetnames or tuple(c.value for c in wb[sheet][1]) != _V11_SHEETS[sheet]:
        errors.append("ATTRIBUTE_SHEET_SCHEMA_INVALID:Player_Attribute_Changes")
    else:
        keys=set()
        for row in _v11_rows(wb[sheet]):
            key=(row.get("Player_ID"), row.get("From_Snapshot"), row.get("To_Snapshot"), row.get("Attribute_Name"))
            if key in keys: errors.append("ATTRIBUTE_CHANGE_NATURAL_KEY_DUPLICATE")
            keys.add(key)
            try:
                if int(row.get("Delta")) != int(row.get("New_Value")) - int(row.get("Old_Value")): errors.append("ATTRIBUTE_DELTA_INVALID")
                if _v11_date(row.get("From_Snapshot")) >= _v11_date(row.get("To_Snapshot")): errors.append("ATTRIBUTE_CHANGE_ORDER_INVALID")
            except (TypeError, ValueError): errors.append("ATTRIBUTE_DELTA_INVALID"); continue
            schema = "OUTFIELD" if row.get("Attribute_Name") in _OUTFIELD else "GOALKEEPER" if row.get("Attribute_Name") in _GOALKEEPER else None
            history=sorted(k[2] for k in snapshots if k[0] == str(row.get("Player_ID")) and k[1] == schema)
            if not schema or str(row.get("From_Snapshot")) not in history or str(row.get("To_Snapshot")) not in history or history.index(str(row.get("To_Snapshot"))) != history.index(str(row.get("From_Snapshot"))) + 1:
                errors.append("ATTRIBUTE_CHANGE_PREDECESSOR_INVALID")
    return {"pass": not errors, "errors": sorted(set(errors)), "error_count": len(errors)}


def _v11_expected_delta(baseline, candidate, manifest):
    b=load_workbook(baseline,data_only=False); c=load_workbook(candidate,data_only=False)
    unexpected=[]
    allowed_new=set(manifest.get("new_authority_sheets", []))
    for sheet in set(b.sheetnames) | set(c.sheetnames):
        if sheet not in b.sheetnames:
            if sheet not in allowed_new: unexpected.append(f"UNAPPROVED_NEW_SHEET:{sheet}")
            elif c[sheet].max_row != 1: unexpected.append(f"FIXTURE_OR_DATA_LEAK_IN_NEW_SHEET:{sheet}")
            continue
        if sheet not in c.sheetnames: unexpected.append(f"REMOVED_SHEET:{sheet}"); continue
        bh=headers(b[sheet]); ch=headers(c[sheet])
        if sheet != "Nation_Dim":
            if list(b[sheet].values) != list(c[sheet].values): unexpected.append(f"UNEXPECTED_DATA_CHANGE:{sheet}")
            continue
        # Controlled alias is the sole sanctioned Nation_Dim data change.
        br=list(b[sheet].values); cr=list(c[sheet].values)
        if tuple(br[0]) != tuple(cr[0]): unexpected.append("NATION_DIM_SCHEMA_CHANGED"); continue
        extra=cr[1:]
        for r in br[1:]:
            if r in extra: extra.remove(r)
            else: unexpected.append("NATION_DIM_BASELINE_ROW_CHANGED")
        h=ch; valid=[]
        for r in extra:
            if str(r[h["National_Team_ID"]-1]) == "N-0027" and str(r[h["Alias_Name"]-1]) == "英格兰" and str(r[h["Canonical_Display_Name"]-1]) == "英格蘭": valid.append(r)
            else: unexpected.append("UNAPPROVED_NATION_ALIAS")
        alias_expected = bool(manifest.get("controlled_alias", {}).get("added", True))
        if alias_expected and len(valid) != 1: unexpected.append("CONTROLLED_NATION_ALIAS_MISSING_OR_DUPLICATE")
        if not alias_expected and valid: unexpected.append("CONTROLLED_NATION_ALIAS_UNEXPECTED_REPEAT")
    return unexpected



def _v11_national_reconciliation(wb):
    """Report (never overwrite) source-profile versus explicit V1.1 seasonal totals."""
    if "Player_National_Team_Stats" not in wb.sheetnames or "Player_Profile_Snapshots" not in wb.sheetnames:
        return []
    aggregates={}
    for row in _v11_rows(wb["Player_National_Team_Stats"]):
        level=row.get("Team_Level")
        if level not in {"Senior", "U21"}: continue
        key=(str(row.get("Player_ID")), level)
        bucket=aggregates.setdefault(key,{"Apps":0,"Goals":0})
        bucket["Apps"] += int(row.get("Apps") or row.get("Apps_Raw") or 0)
        bucket["Goals"] += int(row.get("Goals") or 0)
    profiles={}
    for row in _v11_rows(wb["Player_Profile_Snapshots"]):
        pid=str(row.get("Player_ID")); snapshot=str(row.get("Snapshot_Date"))
        if pid not in profiles or snapshot > profiles[pid].get("Snapshot_Date",""): profiles[pid]=row
    report=[]
    for (pid,level), totals in sorted(aggregates.items()):
        profile=profiles.get(pid,{})
        caps_field,goals_field=("Senior_NT_Caps","Senior_NT_Goals") if level == "Senior" else ("U21_Caps","U21_Goals")
        declared={"Apps":int(profile.get(caps_field) or 0),"Goals":int(profile.get(goals_field) or 0)}
        report.append({"Player_ID":pid,"Team_Level":level,"seasonal":totals,"source_declared":declared,"classification":"RECONCILED" if totals == declared else "RECONCILIATION_ISSUE"})
    return report

def validate_v11(baseline: Path, candidate: Path, *, expected_manifest: dict) -> dict:
    """Hard-gate V1.1 candidate: governed deltas plus formal data invariants."""
    baseline,candidate=Path(baseline),Path(candidate)
    wb=load_workbook(candidate,data_only=False)
    attrs=validate_v11_attribute_sheets(candidate)
    pids=_v11_player_ids(wb); player_format=[p for p in pids if not re.fullmatch(r"P-\d{4}",p)]
    profile_fk=[]
    if "Player_Profile_Snapshots" in wb.sheetnames:
        for r in _v11_rows(wb["Player_Profile_Snapshots"]):
            if str(r.get("Player_ID")) not in set(pids): profile_fk.append(r.get("Player_ID"))
    nation_errors=[]
    if "Player_National_Team_Stats" in wb.sheetnames:
        rows=_v11_rows(wb["Player_National_Team_Stats"]); nk=set(); nations={str(r[0]) for r in wb["Nation_Dim"].iter_rows(min_row=2,values_only=True)}
        legacy_levels={"成年隊（使用者確認預設）", "成年隊（預設）"}
        for r in rows:
            level=r.get("Team_Level")
            if level not in {"Senior","U21"} | legacy_levels: nation_errors.append("NATIONAL_TEAM_LEVEL_MISSING")
            # V1 retains source-evidence legacy labels/duplicates.  The strict V1.1
            # business key applies only to new explicit Senior/U21 authority rows.
            if level in {"Senior", "U21"}:
                key=(r.get("Player_ID"),r.get("Nation_ID"),level,r.get("Season_ID"),r.get("Country_Raw"),r.get("Context_Club"))
                if key in nk: nation_errors.append("NATIONAL_TEAM_NATURAL_KEY_DUPLICATE")
                nk.add(key)
                if str(r.get("Nation_ID")) not in nations: nation_errors.append("NATION_ID_INVALID")
                if r.get("Source_ID") in (None, ""): nation_errors.append("NATIONAL_SOURCE_ID_MISSING")
    unexpected=_v11_expected_delta(baseline,candidate,expected_manifest)
    formula=formula_errors(candidate); fk=fk_errors(candidate)
    issues=attrs["errors"] + player_format + profile_fk + nation_errors + formula + fk
    reconciliation=_v11_national_reconciliation(wb)
    return {"pass": not issues and not unexpected, "expected_changes": expected_manifest.get("expected_changes", []), "unexpected_changes": unexpected, "unexpected_changes_count":len(unexpected), "attribute_validation":attrs, "identity_errors":player_format + profile_fk, "national_team_errors":sorted(set(nation_errors)), "national_career_reconciliation":reconciliation, "formula_errors":formula, "foreign_key_errors":fk}


# The formal production contract is maintained in a dedicated validator-owned module.
# This public export preserves the established global_validation.validate_v11 API.
from v11_validation import validate_v11 as _formal_validate_v11

def validate_v11(baseline: Path, candidate: Path, *, expected_manifest: dict | None = None, mode: str = "populated") -> dict:
    """Formal V1.1 gate.  ``expected_manifest`` is deliberately ignored.

    The contract is code-owned so a candidate cannot authorize its own changes.
    """
    return _formal_validate_v11(baseline, candidate, mode=mode)


def _dispatch_main():
    import sys
    if "--mode" not in sys.argv:
        return main()
    ap=argparse.ArgumentParser(); ap.add_argument('--mode',choices=['v11'],required=True); ap.add_argument('--baseline',required=True,type=Path); ap.add_argument('--candidate',required=True,type=Path); ap.add_argument('--manifest',required=True,type=Path); ap.add_argument('--report',required=True,type=Path); args=ap.parse_args()
    report=validate_v11(args.baseline,args.candidate,expected_manifest=json.loads(args.manifest.read_text(encoding='utf-8')))
    args.report.write_text(json.dumps(report,ensure_ascii=False,indent=2)+'\n',encoding='utf-8'); print(json.dumps({'pass':report['pass'],'unexpected_changes':report['unexpected_changes_count']},ensure_ascii=False))
    if not report['pass']: raise SystemExit(1)

if __name__=='__main__': _dispatch_main()
