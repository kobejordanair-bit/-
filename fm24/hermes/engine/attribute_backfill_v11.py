"""Formal V1.1 ATTRIBUTE_BACKFILL operation; always stages away from ACTIVE."""
from __future__ import annotations

from collections import defaultdict
from pathlib import Path
from shutil import copy2
import hashlib
import json
import zipfile

from openpyxl import load_workbook

from generic_importer import ImportFailure, dictrow, headers, refresh, source_id
from full_player_import_v1_core import active_baseline
from full_player_import_v11 import (
    FORMAT,
    SECTIONS,
    _allocate_player,
    _attribute_payload,
    _ensure_attribute_sheets,
    _existing_match,
    _same,
    parse_full_v11,
    validate_v11_production,
)
from full_player_import_v11_attributes import rebuild_changes
from identity_refresh import refresh_identity_map
from derived_rebuild import rebuild_derived
from barcelona_honours_sync_v1 import sync_workbook
from governed_identity import resolve_secondary_existing_player
from import_policy import evidence_same

OPERATION_MODE = "ATTRIBUTE_BACKFILL"
_ALLOWED_SECTIONS = {"PLAYER", "PLAYER_ATTRIBUTES"}
_SEMANTIC_PROTECTED_SHEETS = (
    "Player_Profile_Snapshots",
    "Player_League_Career",
    "Player_Career_Summaries",
    "Player_Club_Competition_Stats",
    "Player_Club_Season_Totals",
    "Player_National_Team_Stats",
    "Barcelona_Player_Season_Stats",
    "Barcelona_Player_Career",
    "Barcelona_Player_Season_Honours",
)


def _declared_sections(text: str) -> set[str]:
    return {line.strip() for line in text.splitlines() if line.strip() in SECTIONS}


def _existing_profile_dobs(wb, player_id: str) -> set[str]:
    ws = wb["Player_Profile_Snapshots"]
    h = headers(ws)
    dob_field = "DOB" if "DOB" in h else "Date_of_Birth"
    return {
        str(row[h[dob_field] - 1])
        for row in ws.iter_rows(min_row=2, values_only=True)
        if str(row[h["Player_ID"] - 1]) == player_id and row[h[dob_field] - 1] not in (None, "")
    }


def _resolve_existing_identity(wb, doc) -> tuple[str, dict | None]:
    """Exact-first existing resolution, then candidate-only governed onboarding."""
    try:
        player_id, new_player = _allocate_player(wb, doc.player)
    except ImportFailure as exc:
        # Secondary resolution is available only when *all* supplied tokens failed
        # exact lookup. Ambiguous exact tokens always remain fail-closed.
        if not str(exc).startswith(("POSSIBLE_EXISTING_PLAYER_IDENTITY:", "PLAYER_NOT_FOUND:")):
            raise
        raw = doc.player.get("Raw_Name") or doc.player.get("Name")
        resolved = resolve_secondary_existing_player(wb, raw, doc.player.get("Date_of_Birth"))
        return resolved["player_id"], {"raw_token": raw, **resolved}
    if new_player:
        # ATTRIBUTE_BACKFILL may never allocate this ID.  Try only the governed
        # deterministic candidate stage before returning an explicit failure.
        raw = doc.player.get("Raw_Name") or doc.player.get("Name")
        resolved = resolve_secondary_existing_player(wb, raw, doc.player.get("Date_of_Birth"))
        return resolved["player_id"], {"raw_token": raw, **resolved}
    evidence_dobs = _existing_profile_dobs(wb, player_id)
    if evidence_dobs:
        supplied = doc.player.get("Date_of_Birth")
        if len(evidence_dobs) != 1 or supplied is not None and supplied not in evidence_dobs:
            raise ImportFailure("ATTRIBUTE_BACKFILL_DOB_MISMATCH")
    return player_id, None


def _validate_existing_identity(wb, doc) -> str:
    """Compatibility wrapper for callers that only require the resolved ID."""
    return _resolve_existing_identity(wb, doc)[0]


def _semantic_digest(path: Path, sheets=_SEMANTIC_PROTECTED_SHEETS) -> str:
    wb = load_workbook(path, data_only=False, read_only=True)
    try:
        payload = tuple((sheet, tuple(wb[sheet].values)) for sheet in sheets)
    finally:
        wb.close()
    return hashlib.sha256(repr(payload).encode("utf-8")).hexdigest()


def import_attribute_backfill_v11(
    output: Path,
    text: str,
    metadata: dict,
    registry: Path = Path("/opt/data/FM24_World_Current.json"),
    *,
    input_workbook: Path | None = None,
    operation_mode: str = OPERATION_MODE,
) -> dict:
    """Stage ATTRIBUTE_BACKFILL against a temporary copy and formal-validate it.

    PLAYER is parsed and identity-validated only.  No PLAYER Profile authority is
    planned, inserted, updated, or conflict-compared in this operation mode.
    """
    if operation_mode != OPERATION_MODE:
        raise ImportFailure("ATTRIBUTE_BACKFILL_OPERATION_MODE_REQUIRED")
    if not text.splitlines() or text.splitlines()[0].strip() != FORMAT:
        raise ImportFailure("FORMAT_VERSION_REQUIRED_OR_UNSUPPORTED")
    extra = _declared_sections(text) - _ALLOWED_SECTIONS
    if extra:
        raise ImportFailure(f"ATTRIBUTE_BACKFILL_SECTION_UNSUPPORTED:{','.join(sorted(extra))}")

    canonical, registry_data = active_baseline(registry)
    base = canonical if input_workbook is None else Path(input_workbook)
    if input_workbook is not None:
        if base.resolve() == canonical.resolve():
            raise ImportFailure("INPUT_WORKBOOK_MUST_BE_TEMPORARY_COPY")
        if not base.is_file():
            raise ImportFailure(f"INPUT_WORKBOOK_NOT_FOUND:{base}")

    doc = parse_full_v11(text)
    wb = load_workbook(base, data_only=False)
    try:
        player_id, alias_onboarding = _resolve_existing_identity(wb, doc)
        source_value, new_source = source_id(metadata, wb["Source_Index"])
        attribute_sheet = "Player_Attr_Snap_" + ("O" if doc.attribute_schema == "OUTFIELD" else "G")
        payload = _attribute_payload(doc, player_id, source_value, (metadata.get("verification_status") or "SOURCE_TRANSCRIBED_V1_1"))
        ah = headers(wb[attribute_sheet])
        natural_key = (player_id, doc.attribute_snapshot, doc.attribute_schema)
        matches = [
            row for row in wb[attribute_sheet].iter_rows(min_row=2, values_only=True)
            if (str(row[ah["Player_ID"] - 1]), str(row[ah["Snapshot_Date"] - 1]), str(row[ah["Attribute_Schema"] - 1])) == natural_key
        ]
        inserts = []
        noops = defaultdict(int)
        if matches:
            if len(matches) > 1:
                raise ImportFailure("BUSINESS_KEY_DUPLICATE_IN_WORKBOOK:PLAYER_ATTRIBUTES")
            ah = headers(wb[attribute_sheet])
            if not evidence_same({k:matches[0][v-1] for k,v in ah.items()}, payload):
                raise ImportFailure("SAME_SOURCE_RECORD_CONFLICT:PLAYER_ATTRIBUTES")
            noops["PLAYER_ATTRIBUTES"] += 1
        else:
            inserts.append(("PLAYER_ATTRIBUTES", attribute_sheet, payload))
    finally:
        wb.close()

    output = Path(output)
    if output.resolve() in {canonical.resolve(), base.resolve()}:
        raise ImportFailure('OUTPUT_MUST_BE_ISOLATED_CANDIDATE')
    stage = Path(str(output) + ".staging.xlsx")
    if stage.exists():
        stage.unlink()
    try:
        copy2(base, stage)
        staged = load_workbook(stage, data_only=False)
        _ensure_attribute_sheets(staged)
        if inserts and new_source:
            source = staged["Source_Index"]
            source.append(dictrow(source, {
                "Source_ID": source_value,
                "File_Name": metadata["file_name"],
                "Local_Path": metadata["local_path"],
                "Domain": "FM World Import V1.1 ATTRIBUTE_BACKFILL",
                "Season_Context": doc.player["Snapshot"],
                "Source_Note": "formal ATTRIBUTE_BACKFILL execution provenance",
            }))
            refresh(source)
        if alias_onboarding:
            dim = staged["Player_Dim"]
            # Append-only controlled alias with all required evidence encoded in
            # Source_Note; submitted bytes remain untouched in Alias_Name.
            note = json.dumps({
                "original_raw_token": alias_onboarding["raw_token"],
                "matched_player_id": player_id,
                "normalized_candidate_form": alias_onboarding["normalized_candidate_form"],
                "dob_evidence": alias_onboarding["dob_evidence"],
                "source_metadata": {"file_name": metadata.get("file_name"), "local_path": metadata.get("local_path")},
                "rule_used": alias_onboarding["rule"],
            }, ensure_ascii=False, sort_keys=True)
            dim.append(dictrow(dim, {
                "Player_ID": player_id,
                "Canonical_Display_Name": next(r[headers(dim)["Canonical_Display_Name"]-1] for r in dim.iter_rows(min_row=2, values_only=True) if str(r[0]) == player_id),
                "Alias_Name": alias_onboarding["raw_token"],
                "Alias_Type": "Controlled exact alias",
                "Verification_Status": (metadata.get("verification_status") or "SOURCE_TRANSCRIBED_V1_1") or "Verified",
                "Source_Note": note,
            }))
            refresh(dim)
        for _, sheet, row in inserts:
            target = staged[sheet]
            target.append(dictrow(target, row))
            refresh(target)
        staged.save(stage)
        staged.close()

        identity = refresh_identity_map(stage, reconcile_row_provenance=True)
        derived = rebuild_derived(stage)
        # Attribute backfill never rewrites Profile/Career source facts; it may
        # nevertheless complete already-authoritative idempotent derivations.
        honours = sync_workbook(stage)
        attribute_changes = rebuild_changes(stage)
        with zipfile.ZipFile(stage) as archive:
            bad_member = archive.testzip()
        if bad_member is not None:
            raise ImportFailure(f"XLSX_INTEGRITY_FAILED:{bad_member}")
        load_workbook(stage, read_only=True).close()
        validation = validate_v11_production(stage, base)
        if not validation["pass"]:
            raise ImportFailure(f"V11_FORMAL_VALIDATION_FAILED:{validation}")
        stage.replace(output)
    except Exception:
        if stage.exists():
            stage.unlink()
        raise

    return {
        "operation_mode": OPERATION_MODE,
        "baseline": str(base),
        "canonical": str(canonical),
        "canonical_sha256": registry_data["sha256"],
        "player_id": player_id,
        "new_player": False,
        "source_id": source_value,
        "rows_added": {"PLAYER": 0, "PLAYER_ATTRIBUTES": len(inserts)},
        "noops": dict(noops),
        "attribute_changes": attribute_changes,
        "identity": identity,
        "derived": derived,
        "honours": honours,
        "validation": validation,
        "validation_pass": True,
    }
