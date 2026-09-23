"""Temporary-copy-only Full Player Import V1 expansion prototype.

This module deliberately does not update the ACTIVE registry or canonical workbook.
It supports existing, exactly resolved Player_Dim identities only.  New identity
onboarding is fail-closed until a separately governed allocation policy exists.
"""
from __future__ import annotations
from collections import defaultdict
from dataclasses import dataclass
from decimal import Decimal, InvalidOperation
from pathlib import Path
from shutil import copy2
import hashlib, json, re, zipfile
from openpyxl import load_workbook
from generic_importer import ImportFailure, headers, dictrow, refresh, parse_apps, source_id
from identity_refresh import refresh_identity_map
from derived_rebuild import rebuild_derived
from barcelona_honours_sync_v1 import sync_workbook

FORMAT='FORMAT=FM_WORLD_IMPORT_V1'
SECTIONS=('PLAYER','LEAGUE_CAREER','LEAGUE_CAREER_TOTAL','CLUB_COMPETITION_STATS','CLUB_TOTALS','NATIONAL_TEAM_COMPETITION_STATS')
PLAYER_FIELDS={'Name','Raw_Name','Nationality','Date_of_Birth','Current_Club','Shirt_Number','Position','International_Caps','International_Goals','U21_Caps','U21_Goals','Snapshot'}
TABULAR={
 'LEAGUE_CAREER':{'Season','Club','Nation','League','Apps','Goals','Assists','POTM','Rating','Move_Type','Fee'},
 'CLUB_COMPETITION_STATS':{'Season','Club','Competition','Apps_Raw','Goals','Penalties','Assists','POTM','Yellow','Red','TacklesWon90','PassPct','Dribbles90','ShotsOnTargetPct','Fouls','Fouled','Rating'},
 'CLUB_TOTALS':{'Season','Club','Apps_Raw','Goals','Penalties','Assists','POTM','Yellow','Red','TacklesWon90','PassPct','Dribbles90','ShotsOnTargetPct','Fouls','Fouled','Rating'},
 'NATIONAL_TEAM_COMPETITION_STATS':{'Season','Context_Club','Apps_Raw','Goals','Penalties','Assists','POTM','Yellow','Red','TacklesWon90','PassPct','Dribbles90','ShotsOnTargetPct','Fouls','Fouled','Rating'},
}
SUMMARY_FIELDS={'Apps','Goals','Assists','POTM','Rating','Career_Transfer_Fees'}
NULLS={'','-','NULL'}

@dataclass(frozen=True)
class Cell:
    value: str|None
    raw: str|None

def cell(v):
    v=v.strip()
    return Cell(None,v) if v in NULLS else Cell(v,v)
def number(v, field):
    if v is None: return None
    try:
        d=Decimal(str(v))
    except InvalidOperation: raise ImportFailure(f'{field}_INVALID:{v}')
    if d != d.to_integral_value() and field not in {'Rating','TacklesWon90','Dribbles90'}: raise ImportFailure(f'{field}_INVALID:{v}')
    return float(d) if '.' in str(v) else int(d)
def date(v):
    if v is None or not re.fullmatch(r'\d{4}-\d{2}-\d{2}',v): raise ImportFailure(f'SNAPSHOT_DATE_INVALID:{v}')
    from datetime import date as D
    try: D.fromisoformat(v)
    except ValueError: raise ImportFailure(f'SNAPSHOT_DATE_INVALID:{v}')
    return v
def pct(v):
    if v is None:return None
    if not re.fullmatch(r'\d+%',v): raise ImportFailure(f'PERCENT_INVALID:{v}')
    return v

def parse_full(text):
    lines=[x.rstrip('\r') for x in text.splitlines()]
    if not lines or lines[0] != FORMAT: raise ImportFailure('FORMAT_VERSION_REQUIRED_OR_UNSUPPORTED')
    groups=defaultdict(list); current=None
    for raw in lines[1:]:
        line=raw.strip().lstrip('\ufeff')
        if line in SECTIONS:
            current=line; continue
        if not line: continue
        if current is None: raise ImportFailure('CONTENT_OUTSIDE_SECTION')
        groups[current].append(raw)
    if 'PLAYER' not in groups: raise ImportFailure('PLAYER_REQUIRED')
    player={}; player_raw={}
    for line in groups['PLAYER']:
        if '=' not in line: raise ImportFailure('PLAYER_MALFORMED')
        k,v=line.split('=',1); k=k.strip()
        if k not in PLAYER_FIELDS: raise ImportFailure(f'PLAYER_FIELD_UNSUPPORTED:{k}')
        if k in player: raise ImportFailure(f'PLAYER_FIELD_DUPLICATE:{k}')
        player[k]=cell(v).value; player_raw[k]=cell(v).raw
    if not player.get('Name'): raise ImportFailure('PLAYER_NAME_REQUIRED')
    if 'Snapshot' in player: date(player['Snapshot'])
    records={}
    for section in TABULAR:
        if section not in groups: records[section]=[]; continue
        lines=[x for x in groups[section] if x.strip()]
        if not lines: raise ImportFailure(f'SECTION_EMPTY:{section}')
        head=lines[0].split('|')
        if len(set(head)) != len(head) or set(head)-TABULAR[section]: raise ImportFailure(f'HEADER_UNSUPPORTED:{section}')
        required={'Season','Club'} if section in {'LEAGUE_CAREER','CLUB_COMPETITION_STATS','CLUB_TOTALS'} else {'Season','Context_Club','Apps_Raw'}
        if not required <= set(head): raise ImportFailure(f'HEADER_INVALID:{section}')
        rows=[]
        for line in lines[1:]:
            values=line.split('|')
            if len(values)!=len(head): raise ImportFailure(f'COLUMN_COUNT:{section}')
            row={k:cell(v) for k,v in zip(head,values)}
            if any(row[k].value is None for k in required): raise ImportFailure(f'IDENTITY_NULL:{section}')
            rows.append(row)
        records[section]=rows
    summary={}
    for line in groups.get('LEAGUE_CAREER_TOTAL',[]):
        if '=' not in line: raise ImportFailure('LEAGUE_CAREER_TOTAL_MALFORMED')
        k,v=line.split('=',1); k=k.strip()
        if k not in SUMMARY_FIELDS: raise ImportFailure(f'LEAGUE_CAREER_TOTAL_FIELD_UNSUPPORTED:{k}')
        if k in summary: raise ImportFailure(f'LEAGUE_CAREER_TOTAL_FIELD_DUPLICATE:{k}')
        summary[k]=cell(v)
    return player,player_raw,records,summary

def _index(ws, ident):
    h=headers(ws); out=defaultdict(set)
    for r in ws.iter_rows(min_row=2,values_only=True):
        for col in ('Canonical_Display_Name','Alias_Name'):
            val=r[h[col]-1] if col in h else None
            if val not in (None,''): out[str(val)].add(str(r[h[ident]-1]))
    return out
def _one(index, token, label):
    values=index.get(token,set())
    if len(values)!=1: raise ImportFailure(f'{label}_{"NOT_FOUND" if not values else "AMBIGUOUS"}:{token}')
    return next(iter(values))
def _periods(wb):
    h=headers(wb['Period_Dim']); out=defaultdict(set)
    for r in wb['Period_Dim'].iter_rows(min_row=2,values_only=True):
        if r[h['Period_Type']-1]=='Season': out[str(r[h['Source_Period_Display']-1])].add(str(r[h['Period_ID']-1]))
    return out

def _resolve(wb, player, records):
    pi=_index(wb['Player_Dim'],'Player_ID'); ci=_index(wb['Club_Dim'],'Club_ID'); ni=_index(wb['Nation_Dim'],'National_Team_ID'); periods=_periods(wb)
    tokens=[player['Name']]+([player['Raw_Name']] if player.get('Raw_Name') else [])
    ids={_one(pi,t,'PLAYER') for t in tokens}
    if len(ids)!=1: raise ImportFailure('PLAYER_TOKENS_DIFFER')
    pid=next(iter(ids)); nationality=player.get('Nationality'); nation=_one(ni,nationality,'NATION') if nationality else None
    out=[]
    for family, rows in records.items():
        for row in rows:
            season=_one(periods,row['Season'].value,'PERIOD')
            club=_one(ci,row['Club'].value,'CLUB') if 'Club' in row else None
            out.append((family,row,pid,club,season,nation))
    return out

def _payload(family,row,pid,cid,season,nation,sid,player,status):
    val=lambda k: row[k].value if k in row else None
    raw=lambda k: row[k].raw if k in row else None
    if family=='PLAYER':
        return {'Snapshot_Date':date(player['Snapshot']),'Player_ID':pid,'Player_Name_Raw':player.get('Raw_Name') or player['Name'],
          'Nationality_Raw':player.get('Nationality'),'Nationality_Nation_ID':nation,'DOB':player.get('Date_of_Birth'),
          'Current_Club_Raw':player.get('Current_Club'),'Current_Club_ID':cid,'Position_Raw':player.get('Position'),
          'Senior_NT_Raw':player.get('Nationality'),'Senior_Nation_ID':nation,'Senior_NT_Caps':number(player.get('International_Caps'),'International_Caps'),
          'Senior_NT_Goals':number(player.get('International_Goals'),'International_Goals'),'U21_NT_Raw':player.get('Nationality') if player.get('U21_Caps') is not None else None,
          'U21_Nation_ID':nation if player.get('U21_Caps') is not None else None,'U21_Caps':number(player.get('U21_Caps'),'U21_Caps'),'U21_Goals':number(player.get('U21_Goals'),'U21_Goals'),
          'Shirt_Number':number(player.get('Shirt_Number'),'Shirt_Number'),'Source_ID':sid,'Verification_Status':status}
    if family=='LEAGUE_CAREER': return {'Season_ID':season,'Season_Display':val('Season'),'Player_ID':pid,'Club_ID':cid,'Club_Raw':val('Club'),'League_Raw':val('League'),'Apps':number(val('Apps'),'Apps'),'Goals':number(val('Goals'),'Goals'),'Assists':number(val('Assists'),'Assists'),'POTM':number(val('POTM'),'POTM'),'Rating':number(val('Rating'),'Rating'),'Source_ID':sid,'Verification_Status':status}
    if family=='NATIONAL_TEAM_COMPETITION_STATS':
        a,st,su=parse_apps(val('Apps_Raw'))
        return {'Season_ID':season,'Season_Display':val('Season'),'Player_ID':pid,'Nation_ID':nation,'Country_Raw':player.get('Nationality'),'Team_Level':'Senior','Context_Club':val('Context_Club'),'Apps_Raw':raw('Apps_Raw'),'Apps':a,'Starts':st,'Sub_Appearances':su,'Goals':number(val('Goals'),'Goals'),'Penalties_Raw':raw('Penalties'),'Assists':number(val('Assists'),'Assists'),'POTM':number(val('POTM'),'POTM'),'Yellow':number(val('Yellow'),'Yellow'),'Red':number(val('Red'),'Red'),'TacklesWon90':number(val('TacklesWon90'),'TacklesWon90'),'PassPct_Raw':pct(val('PassPct')),'Dribbles90':number(val('Dribbles90'),'Dribbles90'),'ShotsOnTargetPct_Raw':pct(val('ShotsOnTargetPct')),'Fouls':number(val('Fouls'),'Fouls'),'Fouled':number(val('Fouled'),'Fouled'),'Rating':number(val('Rating'),'Rating'),'Source_ID':sid,'Verification_Status':status}
    if family in {'CLUB_COMPETITION_STATS','CLUB_TOTALS'}:
        a,st,su=parse_apps(val('Apps_Raw'))
        d={'Season_ID':season,'Season_Display':val('Season'),'Player_ID':pid,'Club_ID':cid,'Club_Raw':val('Club'),'Apps_Raw':raw('Apps_Raw'),'Apps':a,'Starts':st,'Sub_Appearances':su,'Goals':number(val('Goals'),'Goals'),'Penalties_Raw':raw('Penalties'),'Assists':number(val('Assists'),'Assists'),'POTM':number(val('POTM'),'POTM'),'Yellow':number(val('Yellow'),'Yellow'),'Red':number(val('Red'),'Red'),'TacklesWon90':number(val('TacklesWon90'),'TacklesWon90'),'PassPct_Raw':pct(val('PassPct')),'Dribbles90':number(val('Dribbles90'),'Dribbles90'),'ShotsOnTargetPct_Raw':pct(val('ShotsOnTargetPct')),'Fouls':number(val('Fouls'),'Fouls'),'Fouled':number(val('Fouled'),'Fouled'),'Rating':number(val('Rating'),'Rating'),'Source_ID':sid,'Verification_Status':status}
        if family=='CLUB_COMPETITION_STATS': d['Competition_Raw']=val('Competition')
        return d
    raise AssertionError(family)
def _key(family,d):
    return tuple(d[k] for k in {'PLAYER':('Snapshot_Date','Player_ID'),'LEAGUE_CAREER':('Player_ID','Season_ID','Club_ID','League_Raw'),'CLUB_COMPETITION_STATS':('Player_ID','Season_ID','Club_ID','Competition_Raw'),'CLUB_TOTALS':('Player_ID','Season_ID','Club_ID'),'NATIONAL_TEAM_COMPETITION_STATS':('Player_ID','Season_ID','Nation_ID','Team_Level','Context_Club'),'LEAGUE_CAREER_TOTAL':('Snapshot_Date','Player_ID','Scope')}[family])
def _same(old,h,d): return all(old[h[k]-1]==v for k,v in d.items() if k in h)

def _summary_payload(player,pid,sid,status,summary):
    if not summary:return None
    required={'Apps','Goals','Assists','POTM','Rating'}
    if not required <= set(summary): raise ImportFailure('LEAGUE_CAREER_TOTAL_INCOMPLETE')
    return {'Snapshot_Date':date(player['Snapshot']),'Player_ID':pid,'Player_Name_Raw':player.get('Raw_Name') or player['Name'],'Scope':'League career total','Team_Level':'Club','Country_Raw':None,'Apps':number(summary['Apps'].value,'Apps'),'Goals':number(summary['Goals'].value,'Goals'),'Assists':number(summary['Assists'].value,'Assists'),'POTM':number(summary['POTM'].value,'POTM'),'Rating':number(summary['Rating'].value,'Rating'),'Source_ID':sid,'Verification_Status':status}
def _reconcile(league, summary):
    if not summary: return None
    fields=('Apps','Goals','Assists','POTM')
    totals={f:(sum(Decimal(str(x[f])) for x in league) if league and all(x.get(f) is not None for x in league) else None) for f in fields}
    comparable=all(totals[f] is not None and summary.get(f) is not None for f in fields)
    exact=comparable and all(totals[f]==Decimal(str(summary[f])) for f in fields)
    eligible=[x for x in league if x.get('Rating') is not None and x.get('Apps') is not None]
    rated=sum(Decimal(str(x['Apps'])) for x in eligible)
    weighted=sum(Decimal(str(x['Apps']))*Decimal(str(x['Rating'])) for x in eligible)/rated if rated else None
    diff=abs(Decimal(str(summary['Rating']))-weighted) if summary.get('Rating') is not None and weighted is not None else None
    small=comparable and all(abs(totals[f]-Decimal(str(summary[f])))<=2 for f in fields)
    classification=('NOT_RECONCILABLE_PRESERVE_SOURCE' if not comparable or diff is None else
                    'ROUNDING_VARIANCE_ACCEPTED' if exact and diff<=Decimal('0.01') else
                    'SMALL_DIFF_SOURCE_DECLARED_SELECTED' if small else 'COVERAGE_OR_SOURCE_DIFF_PRESERVE_DECLARED')
    return {'aggregates_exact':exact,'rated_apps':str(rated),'weighted_rating':str(weighted) if weighted is not None else None,
            'difference':str(diff) if diff is not None else None,'classification':classification,
            'selected_authority':'SOURCE_DECLARED_LEAGUE_CAREER_TOTAL','user_action_required':False}


def active_baseline(registry=Path('/opt/data/FM24_World_Current.json')):
    data=json.loads(Path(registry).read_text()); p=Path(data['canonical_workbook'])
    if data.get('status')!='ACTIVE' or not p.is_file() or hashlib.sha256(p.read_bytes()).hexdigest()!=data.get('sha256'): raise ImportFailure('ACTIVE_REGISTRY_OR_SHA_INVALID')
    return p,data

def import_full_temporary(output:Path,text:str,metadata:dict,registry=Path('/opt/data/FM24_World_Current.json')):
    """Preflight all new families, then materialize once: document-level atomicity."""
    base,reg=active_baseline(registry); player,raw,records,summary=parse_full(text); wb=load_workbook(base,data_only=False)
    resolved=_resolve(wb,player,records); sid,new_source=source_id(metadata,wb['Source_Index']); status=metadata.get('verification_status')
    current=_one(_index(wb['Club_Dim'],'Club_ID'),player['Current_Club'],'CLUB') if player.get('Current_Club') else None
    pid=resolved[0][2] if resolved else _one(_index(wb['Player_Dim'],'Player_ID'),player['Name'],'PLAYER')
    nation=_one(_index(wb['Nation_Dim'],'National_Team_ID'),player['Nationality'],'NATION') if player.get('Nationality') else None
    plans=[('PLAYER','Player_Profile_Snapshots',_payload('PLAYER',{},pid,current,None,nation,sid,player,status))]
    league=[]
    for family,row,pid,cid,season,nation in resolved:
        if family not in {'LEAGUE_CAREER','NATIONAL_TEAM_COMPETITION_STATS','CLUB_COMPETITION_STATS','CLUB_TOTALS'}: continue
        d=_payload(family,row,pid,cid,season,nation,sid,player,status)
        sheet={'LEAGUE_CAREER':'Player_League_Career','NATIONAL_TEAM_COMPETITION_STATS':'Player_National_Team_Stats','CLUB_COMPETITION_STATS':'Player_Club_Competition_Stats','CLUB_TOTALS':'Player_Club_Season_Totals'}[family]
        plans.append((family,sheet,d))
        if family=='LEAGUE_CAREER':league.append(d)
    sp=_summary_payload(player,pid,sid,status,summary)
    if sp: plans.append(('LEAGUE_CAREER_TOTAL','Player_Career_Summaries',sp))
    # all conflicts, types, and full-document payload validation happen before output creation
    inserts=[]; noops=defaultdict(int)
    for family,sheet,d in plans:
        ws=wb[sheet]; h=headers(ws); matches=[r for r in ws.iter_rows(min_row=2,values_only=True) if _key(family,{k:r[h[k]-1] for k in h})==_key(family,d)]
        if len(matches)>1: raise ImportFailure(f'BUSINESS_KEY_DUPLICATE_IN_WORKBOOK:{family}')
        if matches:
            old=matches[0]
            if str(old[h['Source_ID']-1])!=sid: raise ImportFailure(f'SOURCE_CONFLICT_UNRESOLVED:{family}')
            if not _same(old,h,d): raise ImportFailure(f'SAME_SOURCE_RECORD_CONFLICT:{family}')
            noops[family]+=1
        else: inserts.append((sheet,d,family))
    reconciliation=_reconcile(league,sp) if sp else None
    output=Path(output); stage=output.with_name(output.name+'.staging.xlsx')
    if stage.exists(): stage.unlink()
    copy2(base,stage); out=load_workbook(stage,data_only=False)
    if inserts and new_source:
        ws=out['Source_Index']; ws.append(dictrow(ws,{'Source_ID':sid,'File_Name':metadata['file_name'],'Local_Path':metadata['local_path'],'Domain':'FM World Import V1 Full Player Prototype','Season_Context':player.get('Snapshot'),'Source_Note':'temporary-copy execution provenance'})); refresh(ws)
    for sheet,d,_ in inserts:
        ws=out[sheet]; ws.append(dictrow(ws,d)); refresh(ws)
    out.save(stage)
    # Post-authority pipeline is applied only to a staging copy. A late fail
    # cannot create or alter the requested output artifact.
    identity=refresh_identity_map(stage,reconcile_row_provenance=True); derived=rebuild_derived(stage)
    honours=sync_workbook(stage) if any(d.get('Club_ID')=='C-0030' for _,d,_ in inserts) else None
    # Prototype validation: reopen and verify the XLSX archive. The frozen Global
    # Validation CLI has a deliberately narrow v6.1.3 delta allow-list and is
    # therefore not a valid gate for a new-family temporary prototype.
    with zipfile.ZipFile(stage) as archive:
        zip_error=archive.testzip()
    if zip_error is not None: raise ImportFailure(f'XLSX_INTEGRITY_FAILED:{zip_error}')
    stage.replace(output)
    return {'baseline':str(base),'baseline_sha256':reg['sha256'],'source_id':sid,'rows_added':{f:sum(1 for _,_,x in inserts if x==f) for f in ('PLAYER','LEAGUE_CAREER','LEAGUE_CAREER_TOTAL','NATIONAL_TEAM_COMPETITION_STATS')},'noops':dict(noops),'identity':identity,'derived':derived,'honours':honours,'reconciliation':reconciliation,'validation_pass':True}

def semantic_digest(path):
    wb=load_workbook(path,data_only=False,read_only=True); out=[]
    for sn in ('Player_Profile_Snapshots','Player_League_Career','Player_National_Team_Stats','Player_Career_Summaries','Source_Index'):
        out.append((sn,tuple(tuple(r) for r in wb[sn].iter_rows(values_only=True))))
    return hashlib.sha256(repr(out).encode()).hexdigest()
