"""Cell/key guards and recomputation for mutable materializations, not sheet waivers."""
from __future__ import annotations
import hashlib
import json
from pathlib import Path
from tempfile import TemporaryDirectory
from openpyxl import load_workbook
from generic_importer import headers
from import_policy import LEDGER, LEDGER_HEADER, POLICY, dump, select_row, source_snapshot, period_fields, equivalent

DERIVED_FIELDS = {
 'Barcelona_Player_Season_Stats': {'Apps','Starts','Goals','Assists','Avg_Rating','MOTM','Goals_Conceded','Clean_Sheets','Yellow','Red','Source_ID','Verification_Status'},
 'Barcelona_Player_Career': {'Barcelona_Seasons','Apps','Goals','Assists','MOTM','Avg_Rating','Goals_Conceded','Clean_Sheets',
    'LaLiga_Team_Championships_While_At_Barcelona','UCL_Team_Championships_While_At_Barcelona',
    'Copa_Team_Championships_While_At_Barcelona','Supercopa_Team_Championships_While_At_Barcelona',
    'UEFA_SuperCup_Team_Championships_While_At_Barcelona','Club_World_Cup_Team_Championships_While_At_Barcelona',
    'FIFPro_Count','Ballon_dOr_Top3_Count','Major_Awards','Derivation_Status','Derivation_Sources'},
 'Barcelona_Player_Season_Honours': {'Apps','Attribution_Policy','LaLiga_Team_Championships_While_At_Barcelona',
    'UCL_Team_Championships_While_At_Barcelona','Copa_Team_Championships_While_At_Barcelona',
    'Supercopa_Team_Championships_While_At_Barcelona','UEFA_SuperCup_Team_Championships_While_At_Barcelona',
    'Authority_Lineage','Source_Reference','Verification_Status'},
}
# This adapter never imports personal awards or Club World Cup title facts.
DERIVED_FIELDS['Barcelona_Player_Career'] -= {'Major_Awards', 'FIFPro_Count', 'Ballon_dOr_Top3_Count', 'Derivation_Sources', 'Club_World_Cup_Team_Championships_While_At_Barcelona'}
KEYS = {'Barcelona_Player_Career': ('Player_ID',),
        'Barcelona_Player_Season_Stats': ('Player_ID','Season_ID'),
        'Barcelona_Player_Season_Honours': ('Player_ID','Season_ID')}

def rows(ws):
    h = headers(ws)
    return [{k:r[v-1] for k,v in h.items()} for r in ws.iter_rows(min_row=2,values_only=True) if any(v is not None for v in r)]

def observations(b,c):
    errors=[]; transitions={}
    if LEDGER not in c: return errors, transitions
    if tuple(x.value for x in c[LEDGER][1]) != LEDGER_HEADER: return ['OBSERVATION_HEADER_INVALID'], {}
    old = rows(b[LEDGER]) if LEDGER in b else []
    new = rows(c[LEDGER])
    if new[:len(old)] != old: errors.append('OBSERVATIONS_NOT_APPEND_ONLY')
    ids=set()
    sources={r['Source_ID'] for r in rows(c['Source_Index'])}
    from full_player_import_v11 import _key
    for r in new:
        if r['Observation_ID'] in ids: errors.append('OBSERVATION_DUPLICATE')
        ids.add(r['Observation_ID'])
    for r in new[len(old):]:
        try:
            record={k:v for k,v in r.items() if k!='Observation_ID'}
            if hashlib.sha256(dump(record).encode()).hexdigest()!=r['Observation_ID'] or r['Policy']!=POLICY or r['Source_ID'] not in sources: raise ValueError()
            a,inc,chosen=[json.loads(r[k]) for k in ('Existing_JSON','Submitted_JSON','Selected_JSON')]
            sheet=r['Target_Sheet']; family=r['Family']
            allowed={'PLAYER':'Player_Profile_Snapshots','LEAGUE_CAREER':'Player_League_Career',
                     'CLUB_COMPETITION_STATS':'Player_Club_Competition_Stats','CLUB_TOTALS':'Player_Club_Season_Totals',
                     'NATIONAL_TEAM_COMPETITION_STATS':'Player_National_Team_Stats','LEAGUE_CAREER_TOTAL':'Player_Career_Summaries'}
            if allowed.get(family)!=sheet: raise ValueError()
            expected, rules=select_row(a,inc,family=family,old_snapshot=source_snapshot(b,a.get('Source_ID')),snapshot=r['Snapshot_Date'])
            if expected != chosen or rules != json.loads(r['Rules_JSON']): raise ValueError()
            if a not in rows(b[sheet]) or _key(family,a)!=_key(family,inc): raise ValueError()
            transitions[(sheet,dump(a))] = chosen
        except (ValueError,KeyError,TypeError): errors.append('OBSERVATION_POLICY_INVALID')
    return errors, transitions

def derived_cell_guard(b,c,sheet):
    errors=[]; fields=DERIVED_FIELDS[sheet]; ks=KEYS[sheet]
    def keyed(w):
        out={}
        for r in rows(w[sheet]):
            key=tuple(r.get(k) for k in ks)
            if any(v is None for v in key): continue # preserved legacy presentation row
            if key in out: errors.append('DERIVED_DUPLICATE_KEY:'+sheet)
            out[key]=r
        return out
    old,new=keyed(b),keyed(c)
    for key,r in old.items():
        if key not in new: errors.append('DERIVED_ROW_REMOVED:'+sheet); continue
        for k,v in r.items():
            if k not in fields and not equivalent(v,new[key].get(k)): errors.append('DERIVED_PROTECTED_FIELD:'+sheet+':'+k)
            if sheet == 'Barcelona_Player_Season_Honours' and k != 'Apps' and not equivalent(v,new[key].get(k)):
                errors.append('EXISTING_HONOUR_FACT_PROTECTED:' + k)
    members={(r.get('Player_ID'),r.get('Season_ID')) for r in rows(c['Player_Club_Season_Totals']) if r.get('Club_ID')=='C-0030'}
    for key,r in new.items():
        if key not in old:
            eligible = key[0] in {p for p,s in members} if len(key)==1 else key in members
            if not eligible: errors.append('DERIVED_UNSUPPORTED_NEW_KEY:'+sheet)
            for k,v in r.items():
                if k not in fields|set(ks)|{'Player_Name','Club_ID'} and v is not None: errors.append('DERIVED_NEW_MANUAL_FIELD:'+sheet+':'+k)
    return errors

def recompute_guard(path):
    from derived_rebuild import rebuild_derived
    from barcelona_honours_sync_v1 import sync_workbook
    from full_player_import_v11_attributes import rebuild_changes
    from identity_refresh import refresh_identity_map
    import shutil
    errors=[]
    with TemporaryDirectory(prefix='fm-derived-check-') as td:
        p=Path(td)/'expected.xlsx'; shutil.copy2(path,p)
        try:
            rebuild_derived(p); sync_workbook(p); rebuild_changes(p)
            a=load_workbook(path); e=load_workbook(p)
            for sheet in [*DERIVED_FIELDS, 'Player_Attribute_Changes']:
                if list(a[sheet].values)!=list(e[sheet].values): errors.append('DERIVED_VALUE_NOT_REPRODUCIBLE:'+sheet)
            a.close();e.close()
        except Exception as exc: errors.append('DERIVED_RECOMPUTE_FAILED:'+type(exc).__name__+':'+str(exc))
    return errors
