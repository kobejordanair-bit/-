"""Screenshot receipts -> checked transcription -> immutable workbook/site release.

The Hermes vision model writes the transcription, never a workbook-editing script.
No credentials, model requests or public deployment are performed by this program.
"""
from __future__ import annotations

import argparse
from contextlib import contextmanager
from datetime import datetime, timezone, date
import hashlib
import json
import math
import os
from pathlib import Path
import re
import shutil
import sys
import tempfile
import uuid

ROOT = Path(__file__).resolve().parent
SITE = ROOT.parent
sys.path.insert(0, str(SITE))
sys.path.insert(0, str(ROOT / 'engine'))
FORMAT = 'FM24_SCREENSHOT_BATCH_V1'


class FlowError(ValueError):
    pass


def sha(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def serialized(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, indent=2, allow_nan=False) + '\n'


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))


def write_json(path, value):
    path = Path(path)
    tmp = path.with_name(path.name + '.' + uuid.uuid4().hex + '.tmp')
    try:
        with tmp.open('w', encoding='utf-8', newline='\n') as f:
            f.write(serialized(value)); f.flush(); os.fsync(f.fileno())
        os.replace(tmp, path)
    finally:
        tmp.unlink(missing_ok=True)


def now():
    return datetime.now(timezone.utc).isoformat()


def inside(directory, name):
    root = Path(directory).resolve()
    p = (root / name).resolve()
    if not p.is_relative_to(root) or p == root:
        raise FlowError('PATH_OUTSIDE_BATCH:' + str(name))
    return p


@contextmanager
def registry_lock(registry):
    """Same lock file as the existing Linux registry promotion engine."""
    path = Path(str(registry) + '.lock')
    with path.open('a+b') as f:
        if os.name == 'nt':
            import msvcrt
            if path.stat().st_size == 0:
                f.write(b'0'); f.flush()
            f.seek(0)
            try:
                msvcrt.locking(f.fileno(), msvcrt.LK_NBLCK, 1)
            except OSError as e:
                raise FlowError('IMPORT_ALREADY_RUNNING') from e
            try:
                yield
            finally:
                f.seek(0); msvcrt.locking(f.fileno(), msvcrt.LK_UNLCK, 1)
        else:
            import fcntl
            try:
                fcntl.flock(f.fileno(), fcntl.LOCK_EX | fcntl.LOCK_NB)
            except BlockingIOError as e:
                raise FlowError('IMPORT_ALREADY_RUNNING') from e
            try:
                yield
            finally:
                fcntl.flock(f.fileno(), fcntl.LOCK_UN)


def active_registry(path):
    active = read_json(path)
    required = ('canonical_workbook', 'sha256', 'version', 'schema_version', 'updated_at')
    if active.get('status') != 'ACTIVE' or any(not active.get(k) for k in required):
        raise FlowError('ACTIVE_REGISTRY_INVALID')
    workbook = Path(active['canonical_workbook'])
    if not workbook.is_absolute():
        raise FlowError('ACTIVE_WORKBOOK_PATH_MUST_BE_ABSOLUTE')
    if sha(workbook) != active['sha256']:
        raise FlowError('ACTIVE_BASELINE_SHA256_MISMATCH')
    return active


def config_at(path):
    config = read_json(path)
    if config.get('format') != 'FM24_SCREENSHOT_CONFIG_V1':
        raise FlowError('CONFIG_FORMAT_INVALID')
    for key in ('registry', 'state'):
        if not Path(config[key]).is_absolute():
            raise FlowError('CONFIG_PATH_MUST_BE_ABSOLUTE:' + key)
    return config


def configure(state, *, registry=None, workbook=None):
    """An explicit workbook initializes a private copy, never rewrites sources/."""
    state = Path(state).resolve()
    if state.exists():
        raise FlowError('STATE_ALREADY_EXISTS')
    if (registry is None) == (workbook is None):
        raise FlowError('CHOOSE_REGISTRY_OR_WORKBOOK')
    if registry:
        registry = Path(registry).resolve()
        active_registry(registry)
    else:
        from archive_exchange import validate_workbook
        workbook = Path(workbook).resolve()
        validate_workbook(workbook.read_bytes())
    state.mkdir(parents=True)
    if workbook:
        original = state / 'baseline.xlsx'
        shutil.copyfile(workbook, original)
        registry = state / 'current.json'
        write_json(registry, dict(canonical_workbook=str(original), sha256=sha(original),
                                 version='local-copy', schema_version='FM_WORLD_WORKBOOK_SCHEMA_V2_OBSERVATION_FACTS',
                                 status='ACTIVE', updated_at=now()))
    config = dict(format='FM24_SCREENSHOT_CONFIG_V1', registry=str(registry), state=str(state))
    write_json(state / 'config.json', config)
    return dict(config=str(state / 'config.json'), **doctor(config))


def doctor(config):
    import openpyxl
    import opencc
    active = active_registry(config['registry'])
    from archive_exchange import validate_workbook
    validate_workbook(Path(active['canonical_workbook']).read_bytes())
    with_workbook = openpyxl.load_workbook(active['canonical_workbook'], read_only=True)
    try:
        required = {'Player_Dim', 'Club_Dim', 'Period_Dim', 'Nation_Dim', 'Workbook_Schema_Metadata',
                    'Import_Observations', 'Player_Club_Season_Totals'}
        if required - set(with_workbook.sheetnames):
            raise FlowError('WORKBOOK_SCHEMA_UNSUPPORTED:' + ','.join(sorted(required - set(with_workbook.sheetnames))))
        from screenshot_provenance import FIELDS, TABLES
        for table in TABLES:
            cols = {c.value for c in with_workbook[table][1]}
            if not FIELDS <= cols:
                raise FlowError('OBSERVATION_FACT_COLUMNS_MISSING:' + table)
        count = len(with_workbook.sheetnames)
    finally:
        with_workbook.close()
    return dict(status='READY', workbook=active['canonical_workbook'], baseline_sha256=active['sha256'],
                sheets=count, python=sys.version.split()[0], openpyxl=openpyxl.__version__,
                vision='HERMES_NATIVE_IMAGE_REVIEW_REQUIRED', public_deployment='NOT_AUTOMATIC')


def prepare(config, inputs, *, kind='image'):
    active = active_registry(config['registry'])
    paths = list(dict.fromkeys(Path(p).resolve() for p in inputs))
    if not paths:
        raise FlowError('NO_INPUT_FILES')
    # Inspect all files before creating a receipt. No extension-only image trust.
    evidence = []
    for p in paths:
        if p.stat().st_size > 40 * 1024 * 1024:
            raise FlowError('SOURCE_OVER_40_MB:' + p.name)
        info = dict(sha256=sha(p), original_name=p.name, kind=kind)
        if kind == 'image':
            from PIL import Image
            with Image.open(p) as im:
                im.verify()
            with Image.open(p) as im:
                if im.format not in ('PNG', 'JPEG', 'WEBP'):
                    raise FlowError('IMAGE_FORMAT_UNSUPPORTED:' + str(im.format))
                info.update(width=im.width, height=im.height)
                suffix = {'PNG': '.png', 'JPEG': '.jpg', 'WEBP': '.webp'}[im.format]
        else:
            suffix = '.txt'
            p.read_text(encoding='utf-8-sig')
        info['path'] = 'evidence/' + info['sha256'] + suffix
        evidence.append((p, info))
    unique = {x['sha256']: (p, x) for p, x in evidence}
    batch = Path(config['state']) / 'inbox' / (datetime.now().strftime('%Y%m%d-%H%M%S') + '-' + uuid.uuid4().hex[:8])
    (batch / 'evidence').mkdir(parents=True)
    for p, info in unique.values():
        shutil.copy2(p, batch / info['path'])
    manifest = dict(format=FORMAT, baseline_sha256=active['sha256'], evidence=[x for _, x in unique.values()],
                    documents=[], ignored=[], created_at=now())
    if kind == 'transcript':
        for i, (_, info) in enumerate(unique.values(), 1):
            name = f'transcription-{i}.txt'
            (batch / name).write_text((batch / info['path']).read_text(encoding='utf-8-sig'), encoding='utf-8')
            manifest['documents'].append(dict(path=name, evidence=[info['sha256']],
                                              review=dict(checked=False, unresolved=[])))
    write_json(batch / 'batch.json', manifest)
    return dict(status='AWAITING_TRANSCRIPTION_REVIEW', batch=str(batch / 'batch.json'),
                evidence=manifest['evidence'], next='Read originals, transcribe only visible values, fill documents/review, then run.')


def preflight_text(text):
    from full_player_import_v11 import parse_full_v11
    from full_player_import_v1_core import TABULAR
    doc = parse_full_v11(text)
    # Attributes are optional; a supplied attribute page must still be complete.
    # Never accept non-null fields which the legacy payload builder discards.
    for rows in doc.records.values():
        for row in rows:
            for key in ('Move_Type', 'Fee', 'Nation'):
                if key in row and row[key].value is not None:
                    raise FlowError('UNMATERIALIZED_SOURCE_FIELD:' + key)
    count_fields = {'Apps', 'Goals', 'Assists', 'POTM', 'Yellow', 'Red', 'Fouls', 'Fouled',
                    'Goals_Conceded', 'Clean_Sheets'}
    for family, rows in doc.records.items():
        for row in rows:
            if family == 'CLUB_COMPETITION_STATS' and ('Competition' not in row or row['Competition'].value is None):
                raise FlowError('COMPETITION_SCOPE_REQUIRED')
            for key, cell in row.items():
                v = cell.value
                if v is None:
                    continue
                if key in count_fields | {'Rating', 'TacklesWon90', 'Dribbles90'}:
                    try:
                        n = float(v)
                    except ValueError:
                        raise FlowError('NUMBER_INVALID:' + key)
                    if not math.isfinite(n) or n < 0 or (key in count_fields and not n.is_integer()):
                        raise FlowError('NUMBER_INVALID:' + key)
                    if key == 'Rating' and not 0 <= n <= 10:
                        raise FlowError('RATING_OUT_OF_RANGE')
                if key in ('PassPct', 'ShotsOnTargetPct') and not (re.fullmatch(r'\d+(?:\.\d+)?%', v) and float(v[:-1]) <= 100):
                    raise FlowError('PERCENT_OUT_OF_RANGE:' + key)
    return doc


def checked_batch(path):
    path = Path(path).resolve()
    batch = read_json(path)
    if batch.get('format') != FORMAT:
        raise FlowError('BATCH_FORMAT_INVALID')
    evidence = {}
    for e in batch.get('evidence', []):
        p = inside(path.parent, e['path'])
        if e.get('kind') not in ('image', 'transcript') or sha(p) != e['sha256']:
            raise FlowError('EVIDENCE_HASH_MISMATCH:' + e['path'])
        if e['sha256'] in evidence:
            raise FlowError('DUPLICATE_EVIDENCE')
        evidence[e['sha256']] = e
    if not evidence:
        raise FlowError('EVIDENCE_REQUIRED')
    used = set()
    documents = []
    seen = set()
    for d in batch.get('documents', []):
        review = d.get('review', {})
        if review.get('checked') is not True or review.get('unresolved') != []:
            raise FlowError('NEEDS_REVIEW:' + d['path'])
        refs = set(d.get('evidence', []))
        if not refs or not refs <= evidence.keys():
            raise FlowError('DOCUMENT_EVIDENCE_MISSING:' + d['path'])
        p = inside(path.parent, d['path'])
        text = p.read_text(encoding='utf-8-sig').replace('\r\n', '\n')
        preflight_text(text)
        digest = hashlib.sha256(text.encode()).hexdigest()
        used |= refs
        if digest not in seen:
            documents.append(dict(path=d['path'], text=text, sha256=digest, evidence=sorted(refs)))
            seen.add(digest)
    for item in batch.get('ignored', []):
        if item.get('sha256') not in evidence or not str(item.get('reason', '')).strip():
            raise FlowError('IGNORED_EVIDENCE_REQUIRES_REASON')
        if item['sha256'] in used:
            raise FlowError('EVIDENCE_BOTH_USED_AND_IGNORED')
        used.add(item['sha256'])
    if used != evidence.keys():
        raise FlowError('UNACCOUNTED_SCREENSHOTS:' + ','.join(sorted(evidence.keys() - used)))
    if not documents:
        raise FlowError('NO_CHECKED_DOCUMENTS')
    material = dict(documents=[{k: d[k] for k in ('sha256', 'evidence')} for d in documents],
                    evidence=sorted(evidence), ignored=batch.get('ignored', []))
    batch_id = hashlib.sha256(serialized(material).encode()).hexdigest()
    return batch, documents, batch_id


def semantic_equal(a, b):
    if sha(a) == sha(b):
        return True
    from openpyxl import load_workbook
    x = load_workbook(a, read_only=True, data_only=False)
    y = load_workbook(b, read_only=True, data_only=False)
    try:
        return x.sheetnames == y.sheetnames and all(list(x[s].values) == list(y[s].values) for s in x.sheetnames)
    finally:
        x.close(); y.close()


def build_release(workbook, destination):
    from archive_exchange import runtime, convert, render_html, dumps
    code = runtime()
    packet = convert(Path(workbook), code['engine'])
    packet['filename'] = 'World_Master.xlsx'
    destination.mkdir()
    shutil.copy2(workbook, destination / 'World_Master.xlsx')
    (destination / 'archive.html').write_text(render_html(packet, code, portable=True), encoding='utf-8')
    (destination / 'fm24-exchange.json').write_text(dumps(packet), encoding='utf-8')
    (destination / 'fm24-data.json').write_text(packet['payloadJSON'], encoding='utf-8')
    sums = {p.name: sha(p) for p in destination.iterdir()}
    write_json(destination / 'SHA256SUMS.json', sums)
    return dict(workbook_sha256=packet['workbookSha256'], sheets=len(packet['manifest']['sheets']), rows=packet['manifest']['rows'])


def run_batch(config, manifest, *, activate=True):
    """One batch is one transaction; every candidate and site is built before activation."""
    from production_importer import import_full_player_into_active
    registry = Path(config['registry'])
    manifest = Path(manifest).resolve()
    batch, documents, batch_id = checked_batch(manifest)
    state = Path(config['state'])
    jobs = state / 'runs'; jobs.mkdir(exist_ok=True)
    if jobs.resolve().is_relative_to(manifest.parent):
        raise FlowError('BATCH_DIRECTORY_MUST_NOT_CONTAIN_RUNS')
    job = jobs / (batch_id[:16] + '-' + uuid.uuid4().hex[:8]); job.mkdir()
    report = dict(status='RUNNING', batch_id=batch_id, started_at=now(), steps=[], activated=False,
                  ignored=batch.get('ignored', []), vision_verified_by='Hermes review declaration; not a measured OCR accuracy score')
    try:
        with registry_lock(registry):
            active = active_registry(registry)
            # Replay of the same activated batch requires the active bytes to still match.
            if active.get('screenshot_batch_id') == batch_id:
                report.update(status='NO_OP', reason='Same batch already active', workbook=active['canonical_workbook'])
                return report
            if batch.get('baseline_sha256') != active['sha256']:
                raise FlowError('STALE_BATCH: rerun prepare against the current master and review conflicts')
            report['baseline_sha256'] = active['sha256']
            baseline = Path(active['canonical_workbook'])
            shutil.copyfile(baseline, job / 'baseline.xlsx')
            if sha(job / 'baseline.xlsx') != active['sha256']:
                raise FlowError('BASELINE_CHANGED_DURING_COPY')
            write_json(job / 'previous-registry.json', active)
            # Freeze bytes used in this transaction, including originals and review declarations.
            shutil.copytree(manifest.parent, job / 'input')
            frozen_batch, frozen_documents, frozen_id = checked_batch(job / 'input' / manifest.name)
            if frozen_id != batch_id or frozen_documents != documents:
                raise FlowError('BATCH_CHANGED_DURING_SNAPSHOT')
            current = job / 'baseline.xlsx'
            for i, doc in enumerate(documents, 1):
                local = dict(active, canonical_workbook=str(current), sha256=sha(current))
                write_json(job / 'staging-registry.json', local)
                txt = job / (doc['sha256'] + '.txt'); txt.write_text(doc['text'], encoding='utf-8')
                output = job / f'candidate-{i}.xlsx'
                result = import_full_player_into_active(output, doc['text'],
                    dict(file_name=txt.name, local_path=str(txt), message_id=doc['sha256'][:20],
                         attachment_sequence=1, verification_status='SCREENSHOT_TRANSCRIBED_REVIEWED'),
                    job / 'staging-registry.json')
                gate = result.get('validation', {})
                if (gate.get('pass') is not True or gate.get('unexpected_changes_count') != 0 or
                    gate.get('baseline_sha256') != sha(current) or gate.get('candidate_sha256') != sha(output)):
                    raise FlowError('CANDIDATE_GATE_FAILED')
                write_json(job / f'validation-{i}.json', result)
                report['steps'].append(dict(document=doc['path'], player_id=result['player_id'],
                                           rows_added=result['rows_added'], selected_rows_updated=result['selected_rows_updated'],
                                           observations=result['observations'], validation_mode=gate.get('mode'), candidate_sha256=sha(output)))
                current = output
            report['review_items'] = review_items(job / 'baseline.xlsx', current)
            if semantic_equal(baseline, current):
                report.update(status='NO_OP', reason='No workbook values changed', workbook=str(baseline))
                return report
            report['artifacts'] = build_release(current, job / 'release')
            if sha(current) != report['steps'][-1]['candidate_sha256']:
                raise FlowError('CANDIDATE_CHANGED_AFTER_VALIDATION')
            if sha(job / 'release' / 'World_Master.xlsx') != sha(current):
                raise FlowError('RELEASE_WORKBOOK_HASH_MISMATCH')
            checksums = job / 'release' / 'SHA256SUMS.json'
            if checksums.exists():
                for name, digest in read_json(checksums).items():
                    if sha(inside(job / 'release', name)) != digest:
                        raise FlowError('RELEASE_ARTIFACT_HASH_MISMATCH:' + name)
            # The current registry must still be exactly the registry that was read.
            if active_registry(registry) != active:
                raise FlowError('STALE_CANDIDATE_BASELINE')
            report.update(status='VALIDATED', release=str(job / 'release'), completed_at=now())
            if activate:
                next_active = dict(active, canonical_workbook=str(job / 'release' / 'World_Master.xlsx'),
                                   sha256=report['artifacts']['workbook_sha256'], updated_at=now(),
                                   screenshot_batch_id=batch_id, archive_directory=str(job / 'release'),
                                   archive_sha256=sha(job / 'release' / 'archive.html'),
                                   previous_registry=str(job / 'previous-registry.json'))
                write_json(registry, next_active)
                report.update(status='ACTIVE', activated=True, workbook=next_active['canonical_workbook'])
            return report
    except Exception as exc:
        report.update(status='FAILED', error=str(exc), completed_at=now())
        raise
    finally:
        report.setdefault('completed_at', now())
        try:
            write_json(job / 'result.json', report)
        except OSError as exc:
            if not report.get('activated'):
                raise
            # A successful registry swap must never be described as a failed import
            # merely because the secondary report could not be persisted.
            report['report_write_error'] = str(exc)


def review_items(baseline, candidate):
    """Surface retained source conflicts instead of reporting every import as clean."""
    from openpyxl import load_workbook
    known = set()
    for path, previous in ((baseline, True), (candidate, False)):
        wb = load_workbook(path, read_only=True)
        try:
            ws = wb['Import_Observations']
            rows = ws.iter_rows(values_only=True); cols = next(rows)
            items = []
            for values in rows:
                row = dict(zip(cols, values))
                if previous:
                    known.add(row['Observation_ID']); continue
                if row['Observation_ID'] in known:
                    continue
                rules = json.loads(row['Rules_JSON'])
                unresolved = {k: v for k, v in rules.items() if v.startswith(('SMALL_', 'LARGE_', 'NON_COUNTER_', 'APPS_GROUP_'))}
                if unresolved:
                    items.append(dict(sheet=row['Target_Sheet'], observation=row['Observation_ID'],
                                      rules=unresolved, adopted='Previous conflicting fields retained'))
        finally:
            wb.close()
    return items


def rollback(config, expected_sha):
    registry = Path(config['registry'])
    with registry_lock(registry):
        active = active_registry(registry)
        if active['sha256'] != expected_sha:
            raise FlowError('ROLLBACK_BASELINE_CHANGED')
        previous = active_registry(active['previous_registry'])
        write_json(registry, previous)
        return dict(status='ROLLED_BACK', workbook=previous['canonical_workbook'], sha256=previous['sha256'])


def main():
    for stream in (sys.stdout, sys.stderr):
        if hasattr(stream, 'reconfigure'):
            stream.reconfigure(encoding='utf-8')
    p = argparse.ArgumentParser(description='FM24 截圖更新流程：保留原圖、核對轉錄、主檔與網站一起更新')
    sub = p.add_subparsers(dest='command', required=True)
    init = sub.add_parser('configure'); init.add_argument('--state', type=Path, required=True)
    group = init.add_mutually_exclusive_group(required=True)
    group.add_argument('--registry', type=Path); group.add_argument('--workbook', type=Path)
    for name in ('doctor', 'prepare', 'run', 'rollback'):
        command = sub.add_parser(name); command.add_argument('--config', type=Path, required=True)
        if name == 'prepare':
            command.add_argument('inputs', type=Path, nargs='+'); command.add_argument('--kind', choices=['image','transcript'], default='image')
        elif name == 'run':
            command.add_argument('--batch', type=Path, required=True); command.add_argument('--preview', action='store_true')
        elif name == 'rollback':
            command.add_argument('--expected-sha256', required=True)
    a = p.parse_args()
    try:
        if a.command == 'configure': result = configure(a.state, registry=a.registry, workbook=a.workbook)
        else:
            config = config_at(a.config)
            if a.command == 'doctor': result = doctor(config)
            elif a.command == 'prepare': result = prepare(config, a.inputs, kind=a.kind)
            elif a.command == 'run': result = run_batch(config, a.batch, activate=not a.preview)
            else: result = rollback(config, a.expected_sha256)
        print(serialized(result))
    except Exception as exc:
        print(serialized(dict(status='FAILED', error=str(exc))), file=sys.stderr)
        raise SystemExit(1)


if __name__ == '__main__':
    main()
